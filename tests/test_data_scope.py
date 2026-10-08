from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.db import db_cursor


def login(client, email, password):
    r=client.post("/api/auth/login",json={"email":email,"password":password})
    assert r.status_code==200, r.text


def ids():
    with db_cursor() as cur:
        cur.execute("SELECT id FROM users WHERE email='sales1@crm.local'"); a=cur.fetchone()[0]
        cur.execute("SELECT id FROM users WHERE email='sales2@crm.local'"); b=cur.fetchone()[0]
        cur.execute("SELECT TOP 1 id FROM customers WHERE owner_id=? ORDER BY id",a); ca=cur.fetchone()[0]
        cur.execute("SELECT TOP 1 id FROM customers WHERE owner_id=? ORDER BY id",b); cb=cur.fetchone()[0]
    return a,b,ca,cb


def test_employee_a_cannot_read_employee_b_customer():
    with TestClient(app) as c:
        login(c,'sales1@crm.local','Sales1@123')
        _,_,ca,cb=ids()
        assert c.get(f'/api/customers/{ca}').status_code==200
        r=c.get(f'/api/customers/{cb}')
        assert r.status_code==403
        assert 'ngoài phạm vi dữ liệu' in r.json()['detail']


def test_employee_list_and_export_only_contain_own_data():
    with TestClient(app) as c:
        login(c,'sales1@crm.local','Sales1@123')
        a,_,_,_=ids()
        r=c.get('/api/customers')
        assert r.status_code==200
        assert r.json()['items']
        assert all(x['owner_id']==a for x in r.json()['items'])
        x=c.get('/api/data/export')
        assert x.status_code==200
        wb=load_workbook(BytesIO(x.content),read_only=True)
        ws=wb['KhachHang']
        owners=[row[4].value for row in ws.iter_rows(min_row=2)]
        assert owners and all(name=='Nhân viên 1' for name in owners)


def test_team_leader_sees_group_and_director_sees_all():
    with TestClient(app) as c:
        login(c,'leader@crm.local','Leader@123')
        r=c.get('/api/customers'); assert r.status_code==200
        owners={x['owner_name'] for x in r.json()['items']}
        assert {'Nhân viên 1','Nhân viên 2'}.issubset(owners)
    with TestClient(app) as c:
        login(c,'director@crm.local','Director@123')
        r=c.get('/api/customers'); assert r.status_code==200
        assert len(r.json()['items'])>=4
