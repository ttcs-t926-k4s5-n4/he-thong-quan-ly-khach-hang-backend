"""
SCRUM-79 — Là Nhân viên Marketing, tôi muốn tạo lead thủ công và nhập lead hàng loạt từ Excel,
để đưa danh sách thu được từ hội thảo vào hệ thống ngay hôm sau.

Chức năng:
  * Nhập tay một lead (từ sự kiện, danh thiếp...) — bắt buộc có nguồn, cảnh báo trùng.
  * Tải tệp mẫu Excel.
  * Nhập hàng loạt 2 bước: (1) xem trước + báo lỗi theo từng dòng, (2) xác nhận nhập.
  * Mọi lead nhập vào đều bắt buộc có nguồn (cột "Nguồn" hoặc nguồn mặc định của lô).
"""
import csv
import io
import json
import os
from datetime import date, datetime, timedelta, timezone

from flask import Blueprint, current_app, request, send_file
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from app.core import (
    LEAD_SOURCES,
    LEAD_STATUS_NEW,
    MANUAL_SOURCES,
    ApiError,
    clean_text,
    current_user,
    find_duplicate_lead,
    get_db,
    get_json_body,
    get_lead_or_404,
    insert_lead,
    is_valid_email,
    is_valid_phone,
    lead_to_dict,
    normalize_email,
    normalize_key,
    normalize_phone,
    now_iso,
    resolve_campaign_id,
)

FEATURE = {
    "key": "SCRUM-79",
    "branch": "feature/SCRUM-79-lead-import-excel-backend",
    "name": "Tạo lead thủ công và nhập lead hàng loạt từ Excel",
}

bp = Blueprint("scrum79_lead_import", __name__)

TEMPLATE_HEADERS = ["Họ tên*", "Email", "Số điện thoại", "Công ty", "Nhu cầu quan tâm", "Nguồn", "Chi tiết nguồn", "Ghi chú"]

HEADER_ALIASES = {
    "full_name": ["hoten", "hovaten", "ten", "tenkhachhang", "fullname", "name"],
    "email": ["email", "thudientu", "mail"],
    "phone": ["sodienthoai", "sdt", "dienthoai", "phone", "mobile", "didong"],
    "company": ["congty", "tencongty", "company", "donvi"],
    "interest": ["nhucauquantam", "nhucau", "quantam", "interest"],
    "source": ["nguon", "nguonlead", "source"],
    "source_detail": ["chitietnguon", "tensukien", "sukien", "sourcedetail"],
    "note": ["ghichu", "note"],
}
HEADER_LOOKUP = {alias: field for field, aliases in HEADER_ALIASES.items() for alias in aliases}
SOURCE_LOOKUP = {}
for _code in MANUAL_SOURCES:
    SOURCE_LOOKUP[normalize_key(_code)] = _code
    SOURCE_LOOKUP[normalize_key(LEAD_SOURCES[_code])] = _code
SOURCE_HINT = ", ".join(LEAD_SOURCES[c] for c in MANUAL_SOURCES)
FIELD_LIMITS = {"company": 200, "interest": 2000, "note": 2000, "source_detail": 200}


# ---------------------------------------------------------------------------
# Kiểm tra dữ liệu một lead (dùng chung cho nhập tay và từng dòng Excel)
# ---------------------------------------------------------------------------
def validate_lead_input(raw, default_source="", default_detail=""):
    values, errors = {}, {}
    full_name = clean_text(raw.get("full_name"))
    if not full_name:
        errors["full_name"] = "Họ tên là bắt buộc"
    elif len(full_name) > 150:
        errors["full_name"] = "Họ tên tối đa 150 ký tự"
    values["full_name"] = full_name

    email = normalize_email(raw.get("email"))
    if email and not is_valid_email(email):
        errors["email"] = f"Email '{email}' không hợp lệ"
    values["email"] = email

    phone_raw = clean_text(raw.get("phone"))
    phone = normalize_phone(phone_raw)
    if phone_raw and not is_valid_phone(phone):
        errors["phone"] = f"Số điện thoại '{phone_raw}' không hợp lệ (VD: 0912345678)"
    values["phone"] = phone if phone_raw else ""

    if not email and not phone_raw:
        errors["contact"] = "Cần ít nhất email hoặc số điện thoại"

    for field, limit in FIELD_LIMITS.items():
        value = clean_text(raw.get(field), keep_newlines=field in ("interest", "note"))
        if len(value) > limit:
            errors[field] = f"Tối đa {limit} ký tự"
        values[field] = value

    source_raw = clean_text(raw.get("source")) or clean_text(default_source)
    values["source"] = ""
    if not source_raw:
        errors["source"] = f"Nguồn là bắt buộc. Chọn một trong: {SOURCE_HINT}"
    else:
        code = SOURCE_LOOKUP.get(normalize_key(source_raw))
        if code is None:
            errors["source"] = f"Nguồn '{source_raw}' không hợp lệ. Chọn một trong: {SOURCE_HINT}"
        values["source"] = code or source_raw
    if not values["source_detail"]:
        values["source_detail"] = clean_text(default_detail)
    return values, errors


def _truthy(value):
    return value is True or str(value).strip().lower() in ("1", "true", "yes", "co", "có")


# ---------------------------------------------------------------------------
# Nhập tay một lead
# ---------------------------------------------------------------------------
@bp.post("/api/leads")
def create_lead():
    db = get_db()
    data = get_json_body()
    values, errors = validate_lead_input(data)
    campaign_id, campaign_err = resolve_campaign_id(db, data.get("campaign_id"))
    if campaign_err:
        errors["campaign_id"] = campaign_err
    if errors:
        raise ApiError(422, "Dữ liệu lead chưa hợp lệ", errors)

    duplicate = find_duplicate_lead(db, values["email"], values["phone"])
    if duplicate and not _truthy(data.get("allow_duplicate")):
        raise ApiError(
            409,
            f"Lead đã tồn tại (#{duplicate['id']} - {duplicate['full_name']}). "
            "Gửi allow_duplicate=true nếu vẫn muốn tạo.",
            {"duplicate_lead_id": duplicate["id"]},
        )
    lead_id = insert_lead(db, {**values, "status": LEAD_STATUS_NEW, "campaign_id": campaign_id,
                               "created_by": current_user()})
    return lead_to_dict(get_lead_or_404(db, lead_id)), 201


# ---------------------------------------------------------------------------
# Tệp mẫu
# ---------------------------------------------------------------------------
@bp.get("/api/leads/import/template")
def download_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Danh sách lead"
    ws.append(TEMPLATE_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F6FEB")
    for col, width in zip("ABCDEFGH", (26, 30, 16, 28, 40, 16, 28, 30)):
        ws.column_dimensions[col].width = width
    ws.column_dimensions["C"].number_format = "@"   # giữ số 0 đầu của SĐT
    ws.freeze_panes = "A2"

    guide = wb.create_sheet("Hướng dẫn")
    guide.append(["Hướng dẫn nhập lead"])
    guide["A1"].font = Font(bold=True, size=13)
    for line in [
        "1. Điền dữ liệu từ dòng 2 của sheet 'Danh sách lead', không đổi tên cột.",
        "2. Bắt buộc: Họ tên và ít nhất một trong Email / Số điện thoại.",
        "3. Cột Nguồn: để trống nếu đã chọn nguồn mặc định khi tải tệp lên.",
        "4. Hệ thống sẽ xem trước và báo lỗi theo từng dòng trước khi nhập thật.",
        "",
        "Giá trị hợp lệ cho cột Nguồn:",
    ]:
        guide.append([line])
    for code in MANUAL_SOURCES:
        guide.append([LEAD_SOURCES[code], code])
    guide.column_dimensions["A"].width = 70

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return send_file(
        buffer,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="mau-nhap-lead.xlsx",
    )


# ---------------------------------------------------------------------------
# Đọc tệp
# ---------------------------------------------------------------------------
def cell_to_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def read_table(content, ext):
    if ext == ".csv":
