# -*- coding: utf-8 -*-
"""
[SCRUM-36] NHẬP NGƯỜI DÙNG HÀNG LOẠT TỪ TỆP EXCEL
User story: "Là Quản trị hệ thống, tôi muốn nhập danh sách người dùng hàng loạt
từ tệp Excel, để tạo tài khoản cho cả khối kinh doanh trong vài phút."

Tiêu chí hoàn thành:
  1. Tải được tệp mẫu                                  -> GET  /api/nguoi-dung/tep-mau
  2. Xem trước và báo lỗi theo từng dòng trước khi nhập -> POST /api/nguoi-dung/xem-truoc
  3. Dòng lỗi bị bỏ qua, dòng hợp lệ vẫn được nhập,     -> POST /api/nguoi-dung/nhap-hang-loat
     có báo cáo tổng kết
"""
import io
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill
from flask import Blueprint, request, jsonify, send_file

from cau_hinh import CAC_COT_TEP_MAU, CAC_VAI_TRO_HOP_LE, MAT_KHAU_MAC_DINH
from co_so_du_lieu import ket_noi_csdl, ma_hoa_mat_khau, chuyen_thanh_tu_dien
from kiem_tra_du_lieu import doc_va_kiem_tra_tep_excel

api_nguoi_dung = Blueprint("api_nguoi_dung", __name__)


@api_nguoi_dung.get("/api/nguoi-dung/tep-mau")
def tai_tep_mau():
    """[Tiêu chí 1] Tải được tệp mẫu Excel (kèm trang Hướng dẫn tiếng Việt)."""
    so = openpyxl.Workbook()

    # ----- Trang 1: Danh sách người dùng -----
    trang = so.active
    trang.title = "Danh sách người dùng"
    trang.append(CAC_COT_TEP_MAU)
    for o in trang[1]:
        o.font = Font(bold=True, color="FFFFFF")
        o.fill = PatternFill("solid", fgColor="2F5597")
    trang.append(["Nguyễn Văn An", "an.nguyen@congty.vn", "0901234567",
                  "Nhóm kinh doanh 1", "Nhân viên kinh doanh", "123456"])
    trang.append(["Lê Thị Bình", "binh.le@congty.vn", "+84912345678",
                  "Nhóm kinh doanh 2", "Trưởng nhóm", ""])
    for chu_cai, rong in zip("ABCDEF", [22, 28, 16, 20, 22, 12]):
        trang.column_dimensions[chu_cai].width = rong

    # ----- Trang 2: Hướng dẫn -----
    huong_dan = so.create_sheet("Hướng dẫn")
    cac_dong_huong_dan = [
        "HƯỚNG DẪN ĐIỀN TỆP NHẬP NGƯỜI DÙNG HÀNG LOẠT",
        "",
        "1. Các cột có dấu (*) là bắt buộc: Họ và tên, Email, Vai trò.",
        "2. Email không được trùng với tài khoản đã có trong hệ thống và không trùng nhau trong tệp.",
        "3. Số điện thoại (nếu điền) phải đúng định dạng Việt Nam: bắt đầu bằng 0 hoặc +84, ví dụ 0912345678.",
        "4. Vai trò phải là một trong các giá trị: " + ", ".join(CAC_VAI_TRO_HOP_LE) + ".",
        "5. Mật khẩu có thể bỏ trống, hệ thống sẽ dùng mật khẩu mặc định: " + MAT_KHAU_MAC_DINH + ".",
        "6. Không thay đổi thứ tự hay xóa dòng tiêu đề ở trang 'Danh sách người dùng'.",
        "",
        "Quy trình khuyến nghị: gọi 'Xem trước' để kiểm tra lỗi từng dòng, sửa xong rồi mới 'Nhập hàng loạt'.",
    ]
    for dong in cac_dong_huong_dan:
        huong_dan.append([dong])
    huong_dan["A1"].font = Font(bold=True, size=13)
    huong_dan.column_dimensions["A"].width = 110

    bo_nho = io.BytesIO()
    so.save(bo_nho)
    bo_nho.seek(0)
    return send_file(
        bo_nho,
        as_attachment=True,
        download_name="tep_mau_nhap_nguoi_dung.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@api_nguoi_dung.post("/api/nguoi-dung/xem-truoc")
def xem_truoc_truoc_khi_nhap():
    """[Tiêu chí 2] Xem trước, báo lỗi theo TỪNG DÒNG - chưa ghi vào CSDL."""
    if "tep" not in request.files:
        return jsonify({
            "thanh_cong": False,
            "thong_bao": "Chưa chọn tệp. Hãy gửi form-data với khóa 'tep' chứa tệp Excel.",
        }), 400

    ket_qua, loi_chung = doc_va_kiem_tra_tep_excel(request.files["tep"])
    if loi_chung:
        return jsonify({"thanh_cong": False, "thong_bao": loi_chung}), 400

    so_hop_le = sum(1 for d in ket_qua if d["hop_le"])
    so_loi = len(ket_qua) - so_hop_le
