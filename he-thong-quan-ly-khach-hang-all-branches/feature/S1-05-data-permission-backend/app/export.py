"""Xuất Excel — dùng đúng dữ liệu đã lọc theo phạm vi từ repository."""
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .messages import SCOPE_LABELS
from .resources import OWNER_COLUMNS, Resource

# Dòng tiêu đề cột nằm ở dòng 4 (sau tiêu đề, dòng thông tin, dòng trống)
HEADER_ROW = 4


def build_workbook(resource: Resource, rows, scope_value: str, exported_by: str) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = resource.sheet_title

    columns = resource.columns + OWNER_COLUMNS
    ws.append([f"Danh sách {resource.label}"])
    ws.append([f"Phạm vi: {SCOPE_LABELS[scope_value]} | Người xuất: {exported_by} | "
               f"Thời gian: {datetime.now():%d/%m/%Y %H:%M}"])
    ws.append([])
    ws.append([title for _, title in columns])
    for cell in ws[HEADER_ROW]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    ws["A1"].font = Font(bold=True, size=14)

    for row in rows:
        ws.append([row.get(key) for key, _ in columns])

    for idx, (key, title) in enumerate(columns, start=1):
        width = max([len(str(title))] + [len(str(r.get(key) or "")) for r in rows])
        ws.column_dimensions[get_column_letter(idx)].width = min(width + 2, 50)
    ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=1)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
