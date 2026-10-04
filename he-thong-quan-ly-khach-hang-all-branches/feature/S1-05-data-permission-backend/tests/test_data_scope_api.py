"""Kiểm thử tự động cho tiêu chí chấp nhận SCRUM-29."""
import unittest
from io import BytesIO

from openpyxl import load_workbook

from app.export import HEADER_ROW
from tests.base import (GIAM_DOC, KH_CUA_A, KH_CUA_B, KH_NHOM_BAC, KH_NHOM_NAM,
                        NHAN_VIEN_A, NHAN_VIEN_B, NHAN_VIEN_C, RESOURCE_NAMES,
                        TAT_CA_KH, TRUONG_NHOM_BAC, TRUONG_NHOM_NAM, ApiTestCase)


class NhanVienAKhongDocDuocKhachHangCuaNhanVienB(ApiTestCase):
    """Tiêu chí 4: chứng minh nhân viên A không đọc được khách hàng của nhân viên B."""

    def test_danh_sach_cua_A_khong_chua_khach_hang_cua_B(self):
        ids = self.ids("/api/customers", NHAN_VIEN_A)
        self.assertEqual(ids, KH_CUA_A)
        self.assertTrue(ids.isdisjoint(KH_CUA_B))

    def test_A_mo_chi_tiet_khach_hang_cua_B_bi_chan(self):
        for cid in KH_CUA_B:
            res = self.get(f"/api/customers/{cid}", NHAN_VIEN_A)
            self.assertEqual(res.status_code, 403)
            body = res.get_json()
            self.assertEqual(body["error"], "NGOAI_PHAM_VI_DU_LIEU")
            self.assertNotIn("name", body)  # không lộ dữ liệu

    def test_A_tim_kiem_ten_khach_hang_cua_B_khong_ra_ket_qua(self):
        res = self.get("/api/customers?q=Hòa Bình", NHAN_VIEN_A)
        self.assertEqual(res.get_json()["total"], 0)
        # Đối chứng: chính B tìm thì thấy
        res_b = self.get("/api/customers?q=Hòa Bình", NHAN_VIEN_B)
        self.assertEqual({i["id"] for i in res_b.get_json()["items"]}, {3})

    def test_A_xuat_excel_khong_co_khach_hang_cua_B(self):
        res = self.get("/api/customers/export", NHAN_VIEN_A)
        self.assertEqual(res.status_code, 200)
        self.assertEqual(self._excel_ids(res.data), KH_CUA_A)

    def test_A_khong_the_tu_mo_rong_pham_vi(self):
        for scope in ("team", "all"):
            res = self.get(f"/api/customers?scope={scope}", NHAN_VIEN_A)
            self.assertEqual(res.status_code, 403)
            self.assertEqual(res.get_json()["error"], "VUOT_PHAM_VI_VAI_TRO")
            res = self.get(f"/api/customers/export?scope={scope}", NHAN_VIEN_A)
            self.assertEqual(res.status_code, 403)

    def test_A_khong_doc_duoc_du_lieu_cua_B_o_ca_4_danh_muc(self):
        for name in RESOURCE_NAMES:
            with self.subTest(resource=name):
                items = self.get(f"/api/{name}", NHAN_VIEN_A).get_json()["items"]
                self.assertTrue(items)
                self.assertTrue(all(i["owner_id"] == NHAN_VIEN_A for i in items))
                for rid in KH_CUA_B:  # dữ liệu mẫu: id cơ hội/hoạt động/báo giá = id KH
                    self.assertEqual(self.get(f"/api/{name}/{rid}", NHAN_VIEN_A).status_code, 403)

    @staticmethod
    def _excel_ids(data: bytes) -> set:
        ws = load_workbook(BytesIO(data)).active
        return {row[0] for row in ws.iter_rows(min_row=HEADER_ROW + 1, values_only=True)
                if row[0] is not None}


class BaPhamViDuLieu(ApiTestCase):
    """Tiêu chí 1: của tôi / của nhóm tôi / tất cả — áp dụng cho 4 danh mục."""

    def test_truong_nhom_thay_toan_nhom_minh(self):
        self.assertEqual(self.ids("/api/customers", TRUONG_NHOM_BAC), KH_NHOM_BAC)
        self.assertEqual(self.ids("/api/customers", TRUONG_NHOM_NAM), KH_NHOM_NAM)

    def test_truong_nhom_khong_thay_nhom_khac(self):
        self.assertTrue(self.ids("/api/customers", TRUONG_NHOM_BAC).isdisjoint(KH_NHOM_NAM))
        self.assertEqual(self.get("/api/customers/5", TRUONG_NHOM_BAC).status_code, 403)

    def test_truong_nhom_chon_cua_toi(self):
        self.assertEqual(self.ids("/api/customers?scope=mine", TRUONG_NHOM_BAC), {6})

    def test_giam_doc_thay_tat_ca(self):
        self.assertEqual(self.ids("/api/customers", GIAM_DOC), TAT_CA_KH)
        for cid in TAT_CA_KH:
            self.assertEqual(self.get(f"/api/customers/{cid}", GIAM_DOC).status_code, 200)

    def test_giam_doc_loc_cua_toi(self):
        self.assertEqual(self.ids("/api/customers?scope=mine", GIAM_DOC), {8})

    def test_ca_4_danh_muc_deu_ap_dung_pham_vi(self):
        for name in RESOURCE_NAMES:
            with self.subTest(resource=name):
                self.assertEqual(self.ids(f"/api/{name}", NHAN_VIEN_C), {5})
                self.assertEqual(self.ids(f"/api/{name}", TRUONG_NHOM_BAC), KH_NHOM_BAC)
                self.assertEqual(self.ids(f"/api/{name}", GIAM_DOC), TAT_CA_KH)


class TimKiemXuatExcelVaPhanTrang(ApiTestCase):
    """Tiêu chí 2: mọi truy vấn danh sách tự lọc, kể cả tìm kiếm & xuất Excel."""

    def test_tim_kiem_cua_truong_nhom_chi_trong_nhom(self):
        # "Công ty" khớp nhiều khách hàng ở cả hai nhóm
        ids = self.ids("/api/customers?q=Công ty", TRUONG_NHOM_BAC)
        self.assertTrue(ids)
        self.assertTrue(ids <= KH_NHOM_BAC)

    def test_tong_so_ban_ghi_dung_theo_pham_vi(self):
        body = self.get("/api/customers?page_size=1", NHAN_VIEN_A).get_json()
        self.assertEqual(body["total"], len(KH_CUA_A))
        self.assertEqual(len(body["items"]), 1)

    def test_xuat_excel_theo_pham_vi_va_tu_khoa(self):
        cases = [
            (GIAM_DOC, "", TAT_CA_KH),
            (TRUONG_NHOM_BAC, "", KH_NHOM_BAC),
            (NHAN_VIEN_B, "", KH_CUA_B),
            (GIAM_DOC, "?scope=team", {8}),  # giám đốc không thuộc nhóm -> chỉ của mình
            (GIAM_DOC, "?q=Mekong", {7}),
        ]
        for user, query, expected in cases:
            with self.subTest(user=user, query=query):
                res = self.get(f"/api/customers/export{query}", user)
                self.assertEqual(res.status_code, 200)
                self.assertEqual(NhanVienAKhongDocDuocKhachHangCuaNhanVienB._excel_ids(res.data),
                                 expected)

    def test_xuat_excel_ca_4_danh_muc(self):
        for name in RESOURCE_NAMES:
            with self.subTest(resource=name):
                res = self.get(f"/api/{name}/export", NHAN_VIEN_A)
                self.assertEqual(res.status_code, 200)
                self.assertEqual(NhanVienAKhongDocDuocKhachHangCuaNhanVienB._excel_ids(res.data),
                                 KH_CUA_A)


class ThongBaoTiengViet(ApiTestCase):
    """Tiêu chí 3: truy cập ngoài phạm vi hiển thị thông báo tiếng Việt rõ ràng."""

    def test_thong_bao_ngoai_pham_vi(self):
        msg = self.get("/api/customers/3", NHAN_VIEN_A).get_json()["message"]
        self.assertIn("Bạn không có quyền xem khách hàng #3", msg)
        self.assertIn("ngoài phạm vi dữ liệu của bạn (Của tôi)", msg)

    def test_thong_bao_dung_ten_danh_muc(self):
        expected = {"opportunities": "cơ hội", "activities": "hoạt động", "quotes": "báo giá"}
        for name, label in expected.items():
            msg = self.get(f"/api/{name}/3", NHAN_VIEN_A).get_json()["message"]
            self.assertIn(f"Bạn không có quyền xem {label} #3", msg)

    def test_thong_bao_vuot_pham_vi_vai_tro(self):
        msg = self.get("/api/customers?scope=all", TRUONG_NHOM_BAC).get_json()["message"]
        self.assertIn("Vai trò Trưởng nhóm chỉ được xem dữ liệu trong phạm vi 'Của nhóm tôi'", msg)

    def test_chua_dang_nhap(self):
        res = self.get("/api/customers")
        self.assertEqual(res.status_code, 401)
        self.assertIn("chưa đăng nhập", res.get_json()["message"])

    def test_khong_ton_tai_tra_404(self):
        res = self.get("/api/customers/999", GIAM_DOC)
        self.assertEqual(res.status_code, 404)
        self.assertIn("Không tìm thấy khách hàng #999", res.get_json()["message"])

    def test_json_giu_nguyen_dau_tieng_viet(self):
        raw = self.get("/api/customers/3", NHAN_VIEN_A).data.decode("utf-8")
        self.assertIn("khách hàng", raw)

    def test_api_me_tra_ve_cac_pham_vi_duoc_chon(self):
        body = self.get("/api/me", TRUONG_NHOM_BAC).get_json()
        self.assertEqual(body["max_scope"], "team")
        self.assertEqual([s["value"] for s in body["available_scopes"]], ["mine", "team"])


if __name__ == "__main__":
    unittest.main()
