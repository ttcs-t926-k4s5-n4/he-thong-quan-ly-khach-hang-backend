import io
import os
from datetime import datetime, timedelta, timezone
from typing import Optional

from fastapi import Cookie, FastAPI, HTTPException, Query, Request, Response
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from pydantic import BaseModel, EmailStr

from .db import connect, db_cursor, init_database, row_to_dict, rows_to_dicts
from .security import hash_password, hash_token, new_token, password_error, verify_password
from .seed import seed_demo_data

SESSION_MINUTES = int(os.getenv("SESSION_IDLE_MINUTES", "15"))
RESET_MINUTES = int(os.getenv("RESET_TOKEN_MINUTES", "30"))
FRONTEND_ORIGIN = os.getenv("FRONTEND_ORIGIN", "http://localhost:5173")
COOKIE_NAME = "crm_session"

app = FastAPI(title="Hệ thống quản lý khách hàng", version="3.0.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=[FRONTEND_ORIGIN, "http://127.0.0.1:5173"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.on_event("startup")
def startup():
    init_database()
    seed_demo_data()


def utcnow():
    return datetime.now(timezone.utc).replace(tzinfo=None)


def get_roles(cur, user_id: int) -> list[str]:
    cur.execute("SELECT r.code FROM user_roles ur JOIN roles r ON r.id=ur.role_id WHERE ur.user_id=? ORDER BY r.id", user_id)
    return [r[0] for r in cur.fetchall()]


def public_user(cur, user_id: int):
    cur.execute("""SELECT u.id,u.full_name,u.email,u.status,u.data_scope,u.business_group_id,u.is_active,g.name group_name
                   FROM users u LEFT JOIN business_groups g ON g.id=u.business_group_id WHERE u.id=?""", user_id)
    u = row_to_dict(cur, cur.fetchone())
    if not u: return None
    u["roles"] = get_roles(cur, user_id)
    return u


def create_session(cur, user_id: int, version: int):
    token = new_token()
    cur.execute("INSERT INTO login_sessions(token_hash,user_id,credentials_version,expires_at) VALUES(?,?,?,?)", hash_token(token), user_id, version, utcnow()+timedelta(minutes=SESSION_MINUTES))
    return token


def current_user(request: Request, required: bool = True):
    token = request.cookies.get(COOKIE_NAME)
    if not token:
        if required: raise HTTPException(401, "Phiên đăng nhập đã hết hạn. Hãy đăng nhập lại.")
        return None
    with db_cursor(commit=True) as cur:
        cur.execute("""SELECT s.token_hash,s.user_id,s.credentials_version,s.expires_at,s.revoked_at,u.credentials_version AS user_version,u.is_active
                       FROM login_sessions s JOIN users u ON u.id=s.user_id WHERE s.token_hash=?""", hash_token(token))
        row = row_to_dict(cur, cur.fetchone())
        if not row or row["revoked_at"] is not None or row["expires_at"] < utcnow() or not row["is_active"] or row["credentials_version"] != row["user_version"]:
            if required: raise HTTPException(401, "Phiên đăng nhập đã hết hạn. Hãy đăng nhập lại.")
            return None
        cur.execute("UPDATE login_sessions SET last_seen_at=?, expires_at=? WHERE token_hash=?", utcnow(), utcnow()+timedelta(minutes=SESSION_MINUTES), hash_token(token))
        return public_user(cur, row["user_id"])


def require_role(user, *roles):
    if not set(user["roles"]).intersection(roles):
        raise HTTPException(403, "Bạn không có quyền thực hiện chức năng này. Hãy liên hệ quản trị viên nếu cần được cấp quyền.")


def menu_for(user):
    roles=set(user["roles"])
    items=[
        {"key":"home","label":"Tổng quan","path":"/"},
        {"key":"customers","label":"Dữ liệu kinh doanh","path":"/customers"},
        {"key":"profile","label":"Hồ sơ cá nhân","path":"/profile"},
        {"key":"crm-customers","label":"Khách hàng doanh nghiệp","path":"/crm-customers"},
        {"key":"customer360","label":"Khách hàng 360°","path":"/customer-360"},
        {"key":"customer-import","label":"Nhập khách hàng Excel","path":"/customer-import"},
        {"key":"customer-filter","label":"Tìm & lọc khách hàng","path":"/customer-filter"},
        {"key":"support","label":"Hỗ trợ & rủi ro","path":"/support"},
        {"key":"care","label":"Chăm sóc định kỳ","path":"/care"},
        {"key":"change-password","label":"Đổi mật khẩu","path":"/change-password"},
    ]
    if roles & {"TEAM_LEADER","DIRECTOR","ADMIN"}:
        items.append({"key":"duplicates","label":"Gộp khách trùng","path":"/duplicates"})
    if roles & {"ADMIN","DIRECTOR"}:
        items.extend([
            {"key":"users","label":"Quản lý người dùng","path":"/users"},
            {"key":"products","label":"Sản phẩm & bảng giá","path":"/products"},
            {"key":"org","label":"Cơ cấu kinh doanh","path":"/org"},
            {"key":"categories","label":"Danh mục bán hàng","path":"/categories"},
            {"key":"pipeline","label":"Pipeline","path":"/pipeline"},
            {"key":"reasons","label":"Lý do thắng/thua","path":"/reasons"},
        ])
    if "ADMIN" in roles:
        items.extend([
            {"key":"user-import","label":"Nhập người dùng Excel","path":"/user-import"},
            {"key":"audit","label":"Nhật ký thay đổi","path":"/audit"},
            {"key":"custom-fields","label":"Trường tùy chỉnh","path":"/custom-fields"},
            {"key":"roles","label":"Vai trò & nhóm","path":"/roles"},
            {"key":"lock","label":"Khóa & bàn giao","path":"/account-lock"},
            {"key":"outbox","label":"Email demo","path":"/outbox"},
        ])
    return items


class LoginIn(BaseModel):
    # Dùng str để chấp nhận các email demo nội bộ dạng @crm.local.
    email: str
    password: str

@app.get("/api/health")
def health(): return {"status":"ok","service":"crm-sprint123"}

@app.post("/api/auth/login")
def login(body: LoginIn, response: Response):
    email = body.email.lower()
    generic = "Email hoặc mật khẩu không đúng. Vui lòng kiểm tra lại thông tin đăng nhập."
    locked_message = "Tài khoản đã bị khóa tạm thời. Vui lòng đăng nhập lại sau 15 phút."
    login_failed = False
    account_locked = False
    token = None
    user = None

    # Không raise HTTPException bên trong db_cursor(commit=True), vì context manager
    # sẽ rollback transaction và làm mất failed_login_attempts/locked_until.
    # Thay vào đó ghi DB trước, để context manager commit, rồi mới trả lỗi.
    with db_cursor(commit=True) as cur:
        cur.execute("SELECT id,password_hash,is_active,failed_login_attempts,locked_until,credentials_version FROM users WHERE email=?", email)
        row = row_to_dict(cur, cur.fetchone())
        now = utcnow()

        # Tài khoản đang bị khóa: không cho đăng nhập kể cả mật khẩu đúng.
        if row and row["locked_until"] and row["locked_until"] > now:
            login_failed = True
            account_locked = True

        else:
            # Hết thời gian khóa thì reset bộ đếm rồi cho phép thử lại.
            if row and row["locked_until"] and row["locked_until"] <= now:
                row["failed_login_attempts"] = 0
                row["locked_until"] = None
                cur.execute("UPDATE users SET failed_login_attempts=0, locked_until=NULL WHERE id=?", row["id"])

            if not row or not row["is_active"] or not verify_password(body.password, row["password_hash"]):
                login_failed = True
                # Chỉ lưu bộ đếm cho tài khoản tồn tại. Các lần 1-4 vẫn dùng thông báo chung.
                if row:
                    attempts = int(row["failed_login_attempts"] or 0) + 1
                    locked_until = now + timedelta(minutes=15) if attempts >= 5 else None
                    cur.execute(
                        "UPDATE users SET failed_login_attempts=?, locked_until=? WHERE id=?",
                        attempts, locked_until, row["id"]
                    )
                    # Ngay lần sai thứ 5, báo rõ tài khoản đã bị khóa 15 phút.
                    if locked_until is not None:
                        account_locked = True
            else:
                # Đăng nhập thành công thì xóa bộ đếm sai.
                cur.execute("UPDATE users SET failed_login_attempts=0, locked_until=NULL WHERE id=?", row["id"])
                token = create_session(cur, row["id"], row["credentials_version"])
                user = public_user(cur, row["id"])

    if login_failed:
        if account_locked:
            raise HTTPException(423, locked_message)
        raise HTTPException(401, generic)

    # Session cookie: không Max-Age -> đóng trình duyệt sẽ mất cookie.
    response.set_cookie(COOKIE_NAME, token, httponly=True, samesite="lax", secure=False, path="/")
    return {"message":"Đăng nhập thành công.","user":user,"menu":menu_for(user)}

@app.get("/api/auth/me")
def me(request: Request):
    user=current_user(request)
    return {"user":user,"menu":menu_for(user),"message":"Phiên hợp lệ và đã được gia hạn."}

@app.post("/api/auth/logout")
def logout(request: Request,response: Response):
    token=request.cookies.get(COOKIE_NAME)
    if token:
        with db_cursor(commit=True) as cur:
            cur.execute("UPDATE login_sessions SET revoked_at=? WHERE token_hash=?", utcnow(), hash_token(token))
    response.delete_cookie(COOKIE_NAME,path="/")
    return {"message":"Đã đăng xuất an toàn và thu hồi phiên trên máy chủ."}

class ForgotIn(BaseModel): email: str
@app.post("/api/auth/forgot-password")
def forgot_password(body: ForgotIn):
    with db_cursor(commit=True) as cur:
        cur.execute("SELECT id FROM users WHERE email=?", body.email.lower())
        row=cur.fetchone()
        if row:
            token=new_token()
            cur.execute("UPDATE password_reset_tokens SET used_at=? WHERE user_id=? AND used_at IS NULL", utcnow(), row[0])
            cur.execute("INSERT INTO password_reset_tokens(user_id,token_hash,expires_at) VALUES(?,?,?)", row[0],hash_token(token),utcnow()+timedelta(minutes=RESET_MINUTES))
            link=f"{FRONTEND_ORIGIN}/reset-password?token={token}"
            cur.execute("INSERT INTO outbox_emails(recipient,subject,body) VALUES(?,?,?)", body.email,"Đặt lại mật khẩu CRM",f"Liên kết có hiệu lực 30 phút và chỉ dùng một lần: {link}")
    return {"message":"Nếu email tồn tại trong hệ thống, chúng tôi đã gửi liên kết đặt lại mật khẩu."}

class ResetIn(BaseModel): password: str; password_confirmation: str
@app.post("/api/auth/reset-password/{token}")
def reset_password(token: str, body: ResetIn):
    err=password_error(body.password)
    if err: raise HTTPException(400,err)
    if body.password!=body.password_confirmation: raise HTTPException(400,"Mật khẩu xác nhận không khớp. Hãy nhập lại hai ô giống nhau.")
    with db_cursor(commit=True) as cur:
        cur.execute("SELECT id,user_id,expires_at,used_at FROM password_reset_tokens WHERE token_hash=?",hash_token(token))
        row=row_to_dict(cur,cur.fetchone())
        if not row or row["used_at"] is not None or row["expires_at"]<utcnow():
            raise HTTPException(400,"Liên kết đặt lại mật khẩu không hợp lệ hoặc đã hết hạn. Hãy yêu cầu một liên kết mới.")
        cur.execute("UPDATE users SET password_hash=?,credentials_version=credentials_version+1,updated_at=? WHERE id=?",hash_password(body.password),utcnow(),row["user_id"])
        cur.execute("UPDATE password_reset_tokens SET used_at=? WHERE id=?",utcnow(),row["id"])
        cur.execute("UPDATE login_sessions SET revoked_at=? WHERE user_id=? AND revoked_at IS NULL",utcnow(),row["user_id"])
    return {"message":"Đặt lại mật khẩu thành công. Bạn có thể đăng nhập bằng mật khẩu mới."}

class ChangePasswordIn(BaseModel): current_password:str; new_password:str; new_password_confirmation:str
@app.post("/api/auth/change-password")
def change_password(body:ChangePasswordIn,request:Request):
    user=current_user(request)
    err=password_error(body.new_password)
    if err: raise HTTPException(400,err)
    if body.new_password!=body.new_password_confirmation: raise HTTPException(400,"Xác nhận mật khẩu mới không khớp. Hãy nhập lại.")
    token=request.cookies.get(COOKIE_NAME)
    with db_cursor(commit=True) as cur:
        cur.execute("SELECT password_hash,credentials_version FROM users WHERE id=?",user["id"])
        row=row_to_dict(cur,cur.fetchone())
        if not verify_password(body.current_password,row["password_hash"]): raise HTTPException(400,"Mật khẩu hiện tại không đúng. Hãy kiểm tra lại.")
        if verify_password(body.new_password,row["password_hash"]): raise HTTPException(400,"Mật khẩu mới phải khác mật khẩu hiện tại. Hãy chọn mật khẩu khác.")
        new_ver=row["credentials_version"]+1
        cur.execute("UPDATE users SET password_hash=?,credentials_version=?,updated_at=? WHERE id=?",hash_password(body.new_password),new_ver,utcnow(),user["id"])
        cur.execute("UPDATE login_sessions SET revoked_at=? WHERE user_id=? AND token_hash<>? AND revoked_at IS NULL",utcnow(),user["id"],hash_token(token))
        cur.execute("UPDATE login_sessions SET credentials_version=? WHERE token_hash=?",new_ver,hash_token(token))
    return {"message":"Đổi mật khẩu thành công. Các phiên đăng nhập khác đã được thu hồi."}


def data_scope_clause(user, alias="x"):
    if user["data_scope"]=="all" or "ADMIN" in user["roles"]:
        return "1=1", []
    if user["data_scope"]=="group":
        return f"{alias}.business_group_id=?", [user["business_group_id"]]
    return f"{alias}.owner_id=?", [user["id"]]

@app.get("/api/customers")
def customers(request:Request,q:str=""):
    user=current_user(request); clause,args=data_scope_clause(user,"c")
    sql=f"""SELECT c.id,c.name,c.phone,c.email,c.owner_id,c.business_group_id,u.full_name owner_name,g.name group_name
             FROM customers c JOIN users u ON u.id=c.owner_id LEFT JOIN business_groups g ON g.id=c.business_group_id
             WHERE {clause} AND (c.name LIKE ? OR c.email LIKE ? OR c.phone LIKE ?) ORDER BY c.id DESC"""
    search=f"%{q}%"; args += [search,search,search]
    with db_cursor() as cur:
        cur.execute(sql,*args); return {"items":rows_to_dicts(cur,cur.fetchall()),"scope":user["data_scope"]}

@app.get("/api/opportunities")
def opportunities(request:Request,q:str=""):
    user=current_user(request); clause,args=data_scope_clause(user,"o")
    sql=f"""SELECT o.id,o.title,o.value,o.stage,o.owner_id,o.business_group_id,u.full_name owner_name,g.name group_name
             FROM opportunities o JOIN users u ON u.id=o.owner_id LEFT JOIN business_groups g ON g.id=o.business_group_id
             WHERE {clause} AND (o.title LIKE ? OR o.stage LIKE ?) ORDER BY o.id DESC"""
    search=f"%{q}%"; args += [search,search]
    with db_cursor() as cur:
        cur.execute(sql,*args); return {"items":rows_to_dicts(cur,cur.fetchall()),"scope":user["data_scope"]}

@app.get("/api/activities")
def activities(request:Request,q:str=""):
    user=current_user(request); clause,args=data_scope_clause(user,"a")
    search=f"%{q}%"; args += [search,search]
    with db_cursor() as cur:
        cur.execute(f"""SELECT a.id,a.subject,a.activity_type,a.owner_id,u.full_name owner_name,g.name group_name
                        FROM activities a JOIN users u ON u.id=a.owner_id LEFT JOIN business_groups g ON g.id=a.business_group_id
                        WHERE {clause} AND (a.subject LIKE ? OR a.activity_type LIKE ?) ORDER BY a.id DESC""",*args)
        return {"items":rows_to_dicts(cur,cur.fetchall()),"scope":user["data_scope"]}

@app.get("/api/quotations")
def quotations(request:Request,q:str=""):
    user=current_user(request); clause,args=data_scope_clause(user,"q")
    search=f"%{q}%"; args += [search,search]
    with db_cursor() as cur:
        cur.execute(f"""SELECT q.id,q.quote_no,q.total_amount,q.status,q.owner_id,u.full_name owner_name,g.name group_name
                        FROM quotations q JOIN users u ON u.id=q.owner_id LEFT JOIN business_groups g ON g.id=q.business_group_id
                        WHERE {clause} AND (q.quote_no LIKE ? OR q.status LIKE ?) ORDER BY q.id DESC""",*args)
        return {"items":rows_to_dicts(cur,cur.fetchall()),"scope":user["data_scope"]}



def ensure_record_in_scope(cur, user, table: str, record_id: int):
    allowed = {
        "customers": "c",
        "opportunities": "o",
        "activities": "a",
        "quotations": "q",
    }
    alias = allowed[table]
    cur.execute(f"SELECT {alias}.id FROM {table} {alias} WHERE {alias}.id=?", record_id)
    if not cur.fetchone():
        raise HTTPException(404, "Không tìm thấy bản ghi được yêu cầu.")
    clause, args = data_scope_clause(user, alias)
    cur.execute(f"SELECT {alias}.id FROM {table} {alias} WHERE {alias}.id=? AND {clause}", record_id, *args)
    if not cur.fetchone():
        raise HTTPException(403, "Bạn không có quyền truy cập bản ghi này vì nằm ngoài phạm vi dữ liệu được cấp.")

@app.get("/api/customers/{record_id}")
def customer_detail(record_id:int, request:Request):
    user=current_user(request)
    with db_cursor() as cur:
        ensure_record_in_scope(cur,user,"customers",record_id)
        cur.execute("""SELECT c.id,c.name,c.phone,c.email,c.owner_id,c.business_group_id,u.full_name owner_name,g.name group_name
                       FROM customers c JOIN users u ON u.id=c.owner_id LEFT JOIN business_groups g ON g.id=c.business_group_id WHERE c.id=?""",record_id)
        return {"item":row_to_dict(cur,cur.fetchone()),"scope":user["data_scope"]}

@app.get("/api/opportunities/{record_id}")
def opportunity_detail(record_id:int, request:Request):
    user=current_user(request)
    with db_cursor() as cur:
        ensure_record_in_scope(cur,user,"opportunities",record_id)
        cur.execute("""SELECT o.id,o.title,o.value,o.stage,o.owner_id,o.business_group_id,u.full_name owner_name,g.name group_name
                       FROM opportunities o JOIN users u ON u.id=o.owner_id LEFT JOIN business_groups g ON g.id=o.business_group_id WHERE o.id=?""",record_id)
        return {"item":row_to_dict(cur,cur.fetchone()),"scope":user["data_scope"]}

@app.get("/api/activities/{record_id}")
def activity_detail(record_id:int, request:Request):
    user=current_user(request)
    with db_cursor() as cur:
        ensure_record_in_scope(cur,user,"activities",record_id)
        cur.execute("""SELECT a.id,a.subject,a.activity_type,a.owner_id,a.business_group_id,u.full_name owner_name,g.name group_name
                       FROM activities a JOIN users u ON u.id=a.owner_id LEFT JOIN business_groups g ON g.id=a.business_group_id WHERE a.id=?""",record_id)
        return {"item":row_to_dict(cur,cur.fetchone()),"scope":user["data_scope"]}

@app.get("/api/quotations/{record_id}")
def quotation_detail(record_id:int, request:Request):
    user=current_user(request)
    with db_cursor() as cur:
        ensure_record_in_scope(cur,user,"quotations",record_id)
        cur.execute("""SELECT q.id,q.quote_no,q.total_amount,q.status,q.owner_id,q.business_group_id,u.full_name owner_name,g.name group_name
                       FROM quotations q JOIN users u ON u.id=q.owner_id LEFT JOIN business_groups g ON g.id=q.business_group_id WHERE q.id=?""",record_id)
        return {"item":row_to_dict(cur,cur.fetchone()),"scope":user["data_scope"]}

@app.get("/api/data/export")
def export_scoped_data(request:Request):
    user=current_user(request)
    wb=Workbook(); wb.remove(wb.active)
    specs=[
      ("KhachHang","c","customers c JOIN users u ON u.id=c.owner_id LEFT JOIN business_groups g ON g.id=c.business_group_id",["ID","Tên","Điện thoại","Email","Chủ sở hữu","Nhóm"],"SELECT c.id,c.name,c.phone,c.email,u.full_name,g.name"),
      ("CoHoi","o","opportunities o JOIN users u ON u.id=o.owner_id LEFT JOIN business_groups g ON g.id=o.business_group_id",["ID","Tiêu đề","Giá trị","Giai đoạn","Chủ sở hữu","Nhóm"],"SELECT o.id,o.title,o.value,o.stage,u.full_name,g.name"),
      ("HoatDong","a","activities a JOIN users u ON u.id=a.owner_id LEFT JOIN business_groups g ON g.id=a.business_group_id",["ID","Nội dung","Loại","Chủ sở hữu","Nhóm"],"SELECT a.id,a.subject,a.activity_type,u.full_name,g.name"),
      ("BaoGia","q","quotations q JOIN users u ON u.id=q.owner_id LEFT JOIN business_groups g ON g.id=q.business_group_id",["ID","Số báo giá","Tổng tiền","Trạng thái","Chủ sở hữu","Nhóm"],"SELECT q.id,q.quote_no,q.total_amount,q.status,u.full_name,g.name"),
    ]
    with db_cursor() as cur:
        for title,alias,from_sql,headers,select_sql in specs:
            clause,args=data_scope_clause(user,alias)
            cur.execute(f"{select_sql} FROM {from_sql} WHERE {clause} ORDER BY {alias}.id",*args)
            rows=cur.fetchall(); ws=wb.create_sheet(title); ws.append(headers)
            for row in rows: ws.append(list(row))
    data=io.BytesIO(); wb.save(data); data.seek(0)
    return StreamingResponse(data,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":"attachment; filename=du-lieu-theo-pham-vi.xlsx"})

@app.get("/api/meta")
def meta(request:Request):
    user=current_user(request)
    with db_cursor() as cur:
        cur.execute("SELECT id,code,name FROM roles ORDER BY id"); roles=rows_to_dicts(cur,cur.fetchall())
        cur.execute("SELECT id,name FROM business_groups ORDER BY id"); groups=rows_to_dicts(cur,cur.fetchall())
    return {"roles":roles,"groups":groups,"current_user":user}

@app.get("/api/users")
def list_users(request:Request,q:str="",group:int|None=None,role:str="",status:str="",page:int=1,per_page:int=20):
    user=current_user(request); require_role(user,"ADMIN","DIRECTOR")
    page=max(1,page); per_page=max(1,min(100,per_page)); offset=(page-1)*per_page
    wh=["(u.full_name LIKE ? OR u.email LIKE ? OR g.name LIKE ?)"]; s=f"%{q}%"; params=[s,s,s]
    if group: wh.append("u.business_group_id=?"); params.append(group)
    if status: wh.append("u.status=?"); params.append(status)
    if role: wh.append("EXISTS(SELECT 1 FROM user_roles ur2 JOIN roles r2 ON r2.id=ur2.role_id WHERE ur2.user_id=u.id AND r2.code=?)"); params.append(role)
    where=" AND ".join(wh)
    with db_cursor() as cur:
        cur.execute(f"SELECT COUNT(*) FROM users u LEFT JOIN business_groups g ON g.id=u.business_group_id WHERE {where}",*params); total=cur.fetchone()[0]
        cur.execute(f"""SELECT u.id,u.full_name,u.email,u.status,u.data_scope,u.business_group_id,u.is_active,g.name group_name
                        FROM users u LEFT JOIN business_groups g ON g.id=u.business_group_id WHERE {where}
                        ORDER BY u.id OFFSET ? ROWS FETCH NEXT ? ROWS ONLY""",*(params+[offset,per_page]))
        items=rows_to_dicts(cur,cur.fetchall())
        for item in items:item["roles"]=get_roles(cur,item["id"])
    return {"items":items,"page":page,"per_page":per_page,"total":total,"pages":max(1,(total+per_page-1)//per_page)}

class UserCreateIn(BaseModel):
    full_name:str; email:str; business_group_id:Optional[int]=None; role_ids:list[int]; data_scope:str="self"

def _roles_and_scope(cur, role_ids:list[int], business_group_id:Optional[int]):
    unique_ids=list(dict.fromkeys(role_ids))
    if not unique_ids:
        raise HTTPException(400,"Phải chọn ít nhất một vai trò cho tài khoản.")
    placeholders=",".join("?" for _ in unique_ids)
    cur.execute(f"SELECT id,code FROM roles WHERE id IN ({placeholders})",*unique_ids)
    selected=rows_to_dicts(cur,cur.fetchall())
    if len(selected)!=len(unique_ids):
        raise HTTPException(400,"Danh sách vai trò không hợp lệ. Hãy tải lại dữ liệu.")
    codes={r["code"] for r in selected}
    if "TEAM_LEADER" in codes and not business_group_id:
        raise HTTPException(400,"Trưởng nhóm phải được gán vào một nhóm kinh doanh cụ thể.")
    if business_group_id:
        cur.execute("SELECT 1 FROM business_groups WHERE id=?",business_group_id)
        if not cur.fetchone():
            raise HTTPException(400,"Nhóm kinh doanh không hợp lệ. Hãy chọn lại.")
    if codes & {"ADMIN","DIRECTOR"}:
        scope="all"
    elif "TEAM_LEADER" in codes:
        scope="group"
    else:
        scope="self"
    return unique_ids,codes,scope

@app.post("/api/users")
def create_user_api(body:UserCreateIn,request:Request):
    actor=current_user(request); require_role(actor,"ADMIN")
    temp="Temp@"+new_token()[:8]
    with db_cursor(commit=True) as cur:
        cur.execute("SELECT 1 FROM users WHERE email=?",body.email.lower())
        if cur.fetchone(): raise HTTPException(409,"Email đã tồn tại. Hãy dùng email công ty khác.")
        role_ids,codes,scope=_roles_and_scope(cur,body.role_ids,body.business_group_id)
        cur.execute("INSERT INTO users(full_name,email,password_hash,status,data_scope,business_group_id,is_active) OUTPUT INSERTED.id VALUES(?,?,?,N'pending',?,?,0)",body.full_name,body.email.lower(),hash_password(temp),scope,body.business_group_id)
        uid=cur.fetchone()[0]
        for rid in role_ids:cur.execute("INSERT INTO user_roles(user_id,role_id) VALUES(?,?)",uid,rid)
        token=new_token(); cur.execute("INSERT INTO account_activation_tokens(user_id,token_hash,expires_at) VALUES(?,?,?)",uid,hash_token(token),utcnow()+timedelta(hours=24))
        link=f"{FRONTEND_ORIGIN}/activate?token={token}"
        cur.execute("INSERT INTO outbox_emails(recipient,subject,body) VALUES(?,?,?)",body.email,"Kích hoạt tài khoản CRM",f"Mật khẩu tạm: {temp}\nKích hoạt: {link}")
    return {"message":"Tạo tài khoản thành công. Email kích hoạt kèm mật khẩu tạm đã được tạo trong Email demo.","user_id":uid,"data_scope":scope}

class UserUpdateIn(BaseModel):
    full_name:str; email:str; business_group_id:Optional[int]=None; data_scope:str="self"; status:str="active"
@app.put("/api/users/{user_id}")
def update_user_api(user_id:int,body:UserUpdateIn,request:Request):
    actor=current_user(request);require_role(actor,"ADMIN")
    with db_cursor(commit=True) as cur:
        cur.execute("SELECT id FROM users WHERE email=? AND id<>?",body.email.lower(),user_id)
        if cur.fetchone():raise HTTPException(409,"Email đã tồn tại. Hãy dùng email khác.")
        if body.status not in {"active","pending","locked"}:raise HTTPException(400,"Trạng thái tài khoản không hợp lệ.")
        codes=set(get_roles(cur,user_id))
        if "TEAM_LEADER" in codes and not body.business_group_id:
            raise HTTPException(400,"Trưởng nhóm phải được gán vào một nhóm kinh doanh cụ thể.")
        if body.business_group_id:
            cur.execute("SELECT 1 FROM business_groups WHERE id=?",body.business_group_id)
            if not cur.fetchone():raise HTTPException(400,"Nhóm kinh doanh không hợp lệ. Hãy chọn lại.")
        scope="all" if codes & {"ADMIN","DIRECTOR"} else ("group" if "TEAM_LEADER" in codes else "self")
        is_active=1 if body.status=="active" else 0
        cur.execute("UPDATE users SET full_name=?,email=?,business_group_id=?,data_scope=?,status=?,is_active=?,updated_at=? WHERE id=?",body.full_name,body.email.lower(),body.business_group_id,scope,body.status,is_active,utcnow(),user_id)
        if cur.rowcount==0:raise HTTPException(404,"Không tìm thấy tài khoản. Hãy làm mới danh sách.")
        if not is_active:
            cur.execute("UPDATE login_sessions SET revoked_at=? WHERE user_id=? AND revoked_at IS NULL",utcnow(),user_id)
    return {"message":"Đã cập nhật tài khoản và đồng bộ trạng thái đăng nhập."}

@app.post("/api/auth/activate/{token}")
def activate(token:str):
    with db_cursor(commit=True) as cur:
        cur.execute("SELECT id,user_id,expires_at,used_at FROM account_activation_tokens WHERE token_hash=?",hash_token(token)); row=row_to_dict(cur,cur.fetchone())
        if not row or row["used_at"] or row["expires_at"]<utcnow():raise HTTPException(400,"Liên kết kích hoạt không hợp lệ hoặc hết hạn. Hãy liên hệ quản trị viên.")
        cur.execute("UPDATE users SET status=N'active',is_active=1 WHERE id=?",row["user_id"]);cur.execute("UPDATE account_activation_tokens SET used_at=? WHERE id=?",utcnow(),row["id"])
    return {"message":"Kích hoạt tài khoản thành công."}

class RoleAssignmentIn(BaseModel): role_ids:list[int]; business_group_id:Optional[int]=None
@app.put("/api/users/{user_id}/roles")
def assign_roles(user_id:int,body:RoleAssignmentIn,request:Request):
    actor=current_user(request);require_role(actor,"ADMIN")
    with db_cursor(commit=True) as cur:
        cur.execute("SELECT id FROM users WHERE id=?",user_id)
        if not cur.fetchone():raise HTTPException(404,"Không tìm thấy người dùng cần phân quyền.")
        role_ids,codes,scope=_roles_and_scope(cur,body.role_ids,body.business_group_id)
        if user_id==actor["id"] and "ADMIN" in actor["roles"] and "ADMIN" not in codes:
            raise HTTPException(403,"Bạn không thể tự thu hồi vai trò Quản trị của chính mình. Hãy nhờ một quản trị viên khác thực hiện.")
        cur.execute("DELETE FROM user_roles WHERE user_id=?",user_id)
        for rid in role_ids:cur.execute("INSERT INTO user_roles(user_id,role_id) VALUES(?,?)",user_id,rid)
        cur.execute("UPDATE users SET business_group_id=?,data_scope=?,updated_at=? WHERE id=?",body.business_group_id,scope,utcnow(),user_id)
    return {"message":"Đã cập nhật nhiều vai trò, nhóm kinh doanh và phạm vi dữ liệu tương ứng.","data_scope":scope}

class LockIn(BaseModel): receiver_user_id:int
@app.post("/api/users/{user_id}/lock-and-handover")
def lock_and_handover(user_id:int,body:LockIn,request:Request):
    actor=current_user(request);require_role(actor,"ADMIN")
    if user_id==body.receiver_user_id:raise HTTPException(400,"Người nhận bàn giao phải khác người bị khóa. Hãy chọn nhân viên khác.")
    if user_id==actor["id"]:raise HTTPException(400,"Bạn không thể tự khóa tài khoản đang dùng. Hãy nhờ quản trị viên khác thực hiện.")
    conn=connect();cur=conn.cursor()
    try:
        cur.execute("SELECT id,is_active FROM users WHERE id=?",user_id);target=cur.fetchone()
        cur.execute("SELECT id,is_active FROM users WHERE id=?",body.receiver_user_id);receiver=cur.fetchone()
        if not target:raise HTTPException(404,"Không tìm thấy tài khoản cần khóa.")
        if not target[1]:raise HTTPException(400,"Tài khoản này đã bị khóa trước đó.")
        if not receiver or not receiver[1]:raise HTTPException(400,"Người nhận bàn giao không hợp lệ, chưa kích hoạt hoặc đang bị khóa.")
        cur.execute("SELECT COUNT(*) FROM customers WHERE owner_id=?",user_id);cc=cur.fetchone()[0]
        cur.execute("SELECT COUNT(*) FROM opportunities WHERE owner_id=?",user_id);oc=cur.fetchone()[0]
        cur.execute("UPDATE customers SET owner_id=?,business_group_id=(SELECT business_group_id FROM users WHERE id=?) WHERE owner_id=?",body.receiver_user_id,body.receiver_user_id,user_id)
        cur.execute("UPDATE opportunities SET owner_id=?,business_group_id=(SELECT business_group_id FROM users WHERE id=?) WHERE owner_id=?",body.receiver_user_id,body.receiver_user_id,user_id)
        cur.execute("UPDATE users SET is_active=0,status=N'locked',updated_at=? WHERE id=?",utcnow(),user_id)
        cur.execute("UPDATE login_sessions SET revoked_at=? WHERE user_id=? AND revoked_at IS NULL",utcnow(),user_id)
        cur.execute("INSERT INTO handover_logs(from_user_id,to_user_id,admin_id,customer_count,opportunity_count) VALUES(?,?,?,?,?)",user_id,body.receiver_user_id,actor["id"],cc,oc)
        conn.commit()
    except HTTPException:
        conn.rollback();raise
    except Exception:
        conn.rollback();raise
    finally:conn.close()
    return {"message":"Đã bàn giao toàn bộ khách hàng/cơ hội, ghi nhật ký và khóa tài khoản.","customer_count":cc,"opportunity_count":oc}

@app.get("/api/handover-logs")
def handover_logs(request:Request):
    actor=current_user(request);require_role(actor,"ADMIN")
    with db_cursor() as cur:
        cur.execute("""SELECT TOP 100 h.id,fu.full_name from_user,tu.full_name to_user,au.full_name admin_user,h.customer_count,h.opportunity_count,h.created_at
                       FROM handover_logs h JOIN users fu ON fu.id=h.from_user_id JOIN users tu ON tu.id=h.to_user_id JOIN users au ON au.id=h.admin_id ORDER BY h.id DESC""")
        return {"items":rows_to_dicts(cur,cur.fetchall())}

@app.get("/api/outbox")
def outbox(request:Request):
    actor=current_user(request);require_role(actor,"ADMIN")
    with db_cursor() as cur:
        cur.execute("SELECT TOP 100 id,recipient,subject,body,created_at FROM outbox_emails ORDER BY id DESC")
        return {"items":rows_to_dicts(cur,cur.fetchall())}

# ======================== SPRINT 2 + SPRINT 3 ========================
import json
import re as _re
from pathlib import Path as _Path
from fastapi import File, UploadFile, Form
from fastapi.responses import FileResponse
try:
    from PIL import Image
except Exception:
    Image = None

_UPLOAD_DIR = _Path(__file__).resolve().parents[1] / "uploads"
_UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
_PHONE_VN = _re.compile(r"^(?:0|\+84)(?:3|5|7|8|9)\d{8}$")


def _audit(cur, user_id, action, entity_type, entity_id=None, field_name=None, old_value=None, new_value=None):
    cur.execute("INSERT INTO audit_logs(user_id,action,entity_type,entity_id,field_name,old_value,new_value) VALUES(?,?,?,?,?,?,?)",
                user_id, action, entity_type, entity_id, field_name, None if old_value is None else str(old_value), None if new_value is None else str(new_value))


def _customer_scope(user, alias="c"):
    return data_scope_clause(user, alias)


# ---- S2-01: bulk user import ----
@app.get("/api/s2/users/import-template")
def s2_user_template(request: Request):
    actor=current_user(request); require_role(actor,"ADMIN")
    wb=Workbook(); ws=wb.active; ws.title="Danh sách người dùng"
    ws.append(["Họ và tên","Email","Số điện thoại","Nhóm","Vai trò","Mật khẩu"])
    ws.append(["Nguyễn Văn An","an.nguyen@congty.vn","0912345678","Nhóm Kinh doanh 1","Nhân viên kinh doanh","Sales@123"])
    guide=wb.create_sheet("Hướng dẫn")
    for x in ["Họ tên, Email, Vai trò là bắt buộc.","Email không được trùng.","SĐT Việt Nam: 0xxxxxxxxx hoặc +84xxxxxxxxx.","Vai trò: Quản trị hệ thống, Giám đốc kinh doanh, Trưởng nhóm, Nhân viên kinh doanh."]:
        guide.append([x])
    b=io.BytesIO(); wb.save(b); b.seek(0)
    return StreamingResponse(b,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":"attachment; filename=tep-mau-nguoi-dung.xlsx"})


def _parse_user_xlsx(raw: bytes):
    from openpyxl import load_workbook
    try: wb=load_workbook(io.BytesIO(raw),data_only=True); ws=wb["Danh sách người dùng"] if "Danh sách người dùng" in wb.sheetnames else wb.active
    except Exception: raise HTTPException(400,"Tệp Excel không hợp lệ hoặc bị hỏng.")
    role_map={"Quản trị hệ thống":"ADMIN","Giám đốc kinh doanh":"DIRECTOR","Trưởng nhóm":"TEAM_LEADER","Nhân viên kinh doanh":"SALES"}
    out=[]
    with db_cursor() as cur:
        cur.execute("SELECT LOWER(email) FROM users"); existing={r[0] for r in cur.fetchall()}
        cur.execute("SELECT id,name FROM business_groups"); groups={str(r[1]).strip().lower():r[0] for r in cur.fetchall()}
        cur.execute("SELECT id,code FROM roles"); roles={r[1]:r[0] for r in cur.fetchall()}
    seen=set()
    for idx,row in enumerate(ws.iter_rows(min_row=2,values_only=True),start=2):
        if not any(v is not None and str(v).strip() for v in row): continue
        name=str(row[0] or '').strip(); email=str(row[1] or '').strip().lower(); phone=str(row[2] or '').strip(); group=str(row[3] or '').strip(); role_name=str(row[4] or '').strip(); password=str(row[5] or '').strip() or 'Temp@1234'
        errors=[]
        if not name: errors.append('Thiếu họ tên')
        if '@' not in email: errors.append('Email không hợp lệ')
        if email in existing or email in seen: errors.append('Email bị trùng')
        if phone and not _PHONE_VN.fullmatch(phone): errors.append('Số điện thoại Việt Nam không hợp lệ')
        code=role_map.get(role_name)
        if not code: errors.append('Vai trò không hợp lệ')
        gid=groups.get(group.lower()) if group else None
        if code=='TEAM_LEADER' and not gid: errors.append('Trưởng nhóm phải có nhóm')
        seen.add(email)
        out.append({"row":idx,"valid":not errors,"errors":errors,"data":{"full_name":name,"email":email,"phone":phone,"group":group,"group_id":gid,"role":role_name,"role_code":code,"role_id":roles.get(code) if code else None,"password":password}})
    return out

@app.post("/api/s2/users/import-preview")
async def s2_user_preview(request: Request, file: UploadFile=File(...)):
    actor=current_user(request); require_role(actor,"ADMIN")
    raw=await file.read(); rows=_parse_user_xlsx(raw)
    return {"items":rows,"total":len(rows),"valid":sum(x['valid'] for x in rows),"invalid":sum(not x['valid'] for x in rows)}

@app.post("/api/s2/users/import")
async def s2_user_import(request: Request, file: UploadFile=File(...)):
    actor=current_user(request); require_role(actor,"ADMIN")
    rows=_parse_user_xlsx(await file.read()); created=[]; skipped=[]
    with db_cursor(commit=True) as cur:
        for x in rows:
            if not x['valid']: skipped.append(x); continue
            d=x['data']; scope='all' if d['role_code'] in {'ADMIN','DIRECTOR'} else ('group' if d['role_code']=='TEAM_LEADER' else 'self')
            cur.execute("INSERT INTO users(full_name,email,password_hash,status,data_scope,business_group_id,is_active,phone,email_signature) OUTPUT INSERTED.id VALUES(?,?,?,N'active',?,?,1,?,?)",
                        d['full_name'],d['email'],hash_password(d['password']),scope,d['group_id'],d['phone'],f"Trân trọng,\n{d['full_name']}")
            uid=cur.fetchone()[0]; cur.execute("INSERT INTO user_roles(user_id,role_id) VALUES(?,?)",uid,d['role_id']); created.append({"row":x['row'],"user_id":uid,"email":d['email']})
            _audit(cur,actor['id'],'bulk_import','User',uid,'email',None,d['email'])
    return {"message":f"Đã tạo {len(created)} tài khoản, bỏ qua {len(skipped)} dòng lỗi.","created":created,"skipped":skipped}

# ---- S2-02/03 profile + avatar ----
class ProfileUpdateIn(BaseModel):
    full_name: Optional[str]=None
    phone: Optional[str]=None
    email_signature: Optional[str]=None
    email: Optional[str]=None
    business_group_id: Optional[int]=None
    roles: Optional[list[str]]=None

@app.get("/api/s2/profile")
def s2_profile(request:Request):
    u=current_user(request)
    with db_cursor() as cur:
        cur.execute("SELECT u.id,u.full_name,u.email,u.phone,u.email_signature,u.avatar_path,u.avatar_thumb_path,g.name group_name FROM users u LEFT JOIN business_groups g ON g.id=u.business_group_id WHERE u.id=?",u['id'])
        d=row_to_dict(cur,cur.fetchone()); d['roles']=get_roles(cur,u['id']); return d

@app.put("/api/s2/profile")
def s2_profile_update(body:ProfileUpdateIn,request:Request):
    u=current_user(request)
    if body.email is not None or body.business_group_id is not None or body.roles is not None:
        raise HTTPException(403,"Bạn không được tự thay đổi email, nhóm hoặc vai trò.")
    if body.phone and not _PHONE_VN.fullmatch(body.phone.strip()): raise HTTPException(400,"Số điện thoại không đúng định dạng Việt Nam.")
    with db_cursor(commit=True) as cur:
        cur.execute("SELECT full_name,phone,email_signature FROM users WHERE id=?",u['id']); old=row_to_dict(cur,cur.fetchone())
        name=(body.full_name if body.full_name is not None else old['full_name']).strip(); phone=(body.phone if body.phone is not None else old['phone']); sig=body.email_signature if body.email_signature is not None else old['email_signature']
        if not name: raise HTTPException(400,"Họ tên không được để trống.")
        cur.execute("UPDATE users SET full_name=?,phone=?,email_signature=?,updated_at=? WHERE id=?",name,phone,sig,utcnow(),u['id'])
        for k,nv in [('full_name',name),('phone',phone),('email_signature',sig)]:
            if old.get(k)!=nv: _audit(cur,u['id'],'updated','User',u['id'],k,old.get(k),nv)
    return {"message":"Cập nhật hồ sơ thành công."}

@app.post("/api/s2/profile/avatar")
async def s2_avatar(request:Request, file:UploadFile=File(...)):
    u=current_user(request); raw=await file.read(); ext=(_Path(file.filename or '').suffix or '').lower()
    if ext not in {'.jpg','.jpeg','.png'}: raise HTTPException(400,"Chỉ chấp nhận ảnh JPG hoặc PNG.")
    if len(raw)>2*1024*1024: raise HTTPException(400,"Ảnh vượt quá giới hạn 2MB.")
    if Image is None: raise HTTPException(500,"Máy chủ chưa cài Pillow để xử lý ảnh.")
    try:
        img=Image.open(io.BytesIO(raw)).convert('RGB')
    except Exception: raise HTTPException(400,"Tệp ảnh không hợp lệ.")
    w,h=img.size; side=min(w,h); left=(w-side)//2; top=(h-side)//2; img=img.crop((left,top,left+side,top+side))
    full=_UPLOAD_DIR/f"avatar_{u['id']}.jpg"; thumb=_UPLOAD_DIR/f"avatar_{u['id']}_thumb.jpg"
    img.resize((512,512)).save(full,'JPEG',quality=90); img.resize((128,128)).save(thumb,'JPEG',quality=88)
    with db_cursor(commit=True) as cur:
        cur.execute("UPDATE users SET avatar_path=?,avatar_thumb_path=? WHERE id=?",str(full),str(thumb),u['id']); _audit(cur,u['id'],'updated','User',u['id'],'avatar',None,file.filename)
    return {"message":"Ảnh đã được cắt vuông và tạo bản thu nhỏ 128x128."}

@app.get("/api/s2/profile/avatar/{kind}")
def s2_avatar_get(kind:str,request:Request):
    u=current_user(request)
    with db_cursor() as cur:
        cur.execute("SELECT avatar_path,avatar_thumb_path FROM users WHERE id=?",u['id']); r=cur.fetchone()
    p=(r[1] if kind=='thumb' else r[0]) if r else None
    if not p or not os.path.exists(p): raise HTTPException(404,"Chưa có ảnh đại diện.")
    return FileResponse(p,media_type='image/jpeg')

# ---- S2-04 audit log ----
@app.get("/api/s2/audit-logs")
def s2_audit(request:Request,user_id:int|None=None,entity_type:str='',date_from:str='',date_to:str=''):
    u=current_user(request); require_role(u,'ADMIN')
    wh=['1=1']; args=[]
    if user_id: wh.append('a.user_id=?'); args.append(user_id)
    if entity_type: wh.append('a.entity_type=?'); args.append(entity_type)
    if date_from: wh.append('a.created_at>=?'); args.append(date_from)
    if date_to: wh.append('a.created_at<DATEADD(day,1,?)'); args.append(date_to)
    with db_cursor() as cur:
        cur.execute(f"SELECT TOP 500 a.id,u.full_name user_name,a.action,a.entity_type,a.entity_id,a.field_name,a.old_value,a.new_value,a.created_at FROM audit_logs a LEFT JOIN users u ON u.id=a.user_id WHERE {' AND '.join(wh)} ORDER BY a.id DESC",*args)
        return {"items":rows_to_dicts(cur,cur.fetchall())}

# ---- generic Sprint 2 catalog helpers ----
class ProductIn(BaseModel): code:str; name:str; product_type:str; unit:str; list_price:float; floor_price:float; is_active:bool=True
@app.get('/api/s2/products')
def s2_products(request:Request):
    current_user(request)
    with db_cursor() as cur: cur.execute('SELECT * FROM products ORDER BY id DESC'); return {'items':rows_to_dicts(cur,cur.fetchall())}
@app.post('/api/s2/products')
def s2_product_add(body:ProductIn,request:Request):
    u=current_user(request); require_role(u,'DIRECTOR','ADMIN')
    if body.product_type not in {'Sản phẩm một lần','Dịch vụ thuê bao'}: raise HTTPException(400,'Loại sản phẩm không hợp lệ.')
    if body.floor_price>body.list_price: raise HTTPException(400,'Giá sàn không được cao hơn giá niêm yết.')
    with db_cursor(commit=True) as cur:
        try: cur.execute('INSERT INTO products(code,name,product_type,unit,list_price,floor_price,is_active) OUTPUT INSERTED.id VALUES(?,?,?,?,?,?,?)',body.code,body.name,body.product_type,body.unit,body.list_price,body.floor_price,1 if body.is_active else 0); pid=cur.fetchone()[0]
        except Exception: raise HTTPException(409,'Mã sản phẩm đã tồn tại.')
        _audit(cur,u['id'],'created','Product',pid,'code',None,body.code)
    return {'message':'Đã thêm sản phẩm/dịch vụ.','id':pid}
@app.put('/api/s2/products/{pid}')
def s2_product_update(pid:int,body:ProductIn,request:Request):
    u=current_user(request); require_role(u,'DIRECTOR','ADMIN')
    if body.floor_price>body.list_price: raise HTTPException(400,'Giá sàn không được cao hơn giá niêm yết.')
    with db_cursor(commit=True) as cur:
        cur.execute('SELECT * FROM products WHERE id=?',pid); old=row_to_dict(cur,cur.fetchone())
        if not old: raise HTTPException(404,'Không tìm thấy sản phẩm.')
        cur.execute('UPDATE products SET code=?,name=?,product_type=?,unit=?,list_price=?,floor_price=?,is_active=? WHERE id=?',body.code,body.name,body.product_type,body.unit,body.list_price,body.floor_price,1 if body.is_active else 0,pid)
        for f in ['list_price','floor_price','is_active']:
            if old.get(f)!=getattr(body,f): _audit(cur,u['id'],'updated','Product',pid,f,old.get(f),getattr(body,f))
    return {'message':'Đã cập nhật sản phẩm.'}

class GroupIn(BaseModel): name:str; parent_id:Optional[int]=None; leader_id:Optional[int]=None; region:Optional[str]=None
@app.get('/api/s2/org-groups')
def s2_groups(request:Request):
    current_user(request)
    with db_cursor() as cur:
        cur.execute('SELECT g.id,g.name,g.parent_id,p.name parent_name,g.leader_id,u.full_name leader_name,g.region FROM business_groups g LEFT JOIN business_groups p ON p.id=g.parent_id LEFT JOIN users u ON u.id=g.leader_id ORDER BY g.id'); return {'items':rows_to_dicts(cur,cur.fetchall())}
@app.post('/api/s2/org-groups')
def s2_group_add(body:GroupIn,request:Request):
    u=current_user(request); require_role(u,'DIRECTOR','ADMIN')
    with db_cursor(commit=True) as cur:
        cur.execute('INSERT INTO business_groups(name,parent_id,leader_id,region) OUTPUT INSERTED.id VALUES(?,?,?,?)',body.name,body.parent_id,body.leader_id,body.region); gid=cur.fetchone()[0]
        if body.leader_id: cur.execute("UPDATE users SET business_group_id=?,data_scope='group' WHERE id=?",gid,body.leader_id)
        _audit(cur,u['id'],'created','BusinessGroup',gid,'name',None,body.name)
    return {'message':'Đã tạo nhóm kinh doanh.','id':gid}

class CategoryIn(BaseModel): category_type:str; value:str; sort_order:int=0; is_active:bool=True
@app.get('/api/s2/categories')
def s2_categories(request:Request,category_type:str=''):
    current_user(request)
    with db_cursor() as cur:
        if category_type: cur.execute('SELECT * FROM sales_categories WHERE category_type=? ORDER BY sort_order,id',category_type)
        else: cur.execute('SELECT * FROM sales_categories ORDER BY category_type,sort_order,id')
        return {'items':rows_to_dicts(cur,cur.fetchall())}
@app.post('/api/s2/categories')
def s2_category_add(body:CategoryIn,request:Request):
    u=current_user(request); require_role(u,'DIRECTOR','ADMIN')
    with db_cursor(commit=True) as cur:
        try: cur.execute('INSERT INTO sales_categories(category_type,value,sort_order,is_active) OUTPUT INSERTED.id VALUES(?,?,?,?)',body.category_type,body.value,body.sort_order,1 if body.is_active else 0); cid=cur.fetchone()[0]
        except Exception: raise HTTPException(409,'Giá trị này đang được tham chiếu hoặc đã tồn tại; không thể tạo trùng.')
        _audit(cur,u['id'],'created','SalesCategory',cid,'value',None,body.value)
    return {'message':'Đã thêm danh mục dùng chung.','id':cid}

class CustomFieldIn(BaseModel): entity_type:str; field_name:str; field_type:str; is_required:bool=False; options:list[str]=[]; sort_order:int=0
@app.get('/api/s2/custom-fields')
def s2_custom_fields(request:Request):
    current_user(request)
    with db_cursor() as cur: cur.execute('SELECT * FROM custom_fields WHERE is_active=1 ORDER BY entity_type,sort_order,id'); items=rows_to_dicts(cur,cur.fetchall())
    for x in items:
        try: x['options']=json.loads(x.get('options_json') or '[]')
        except: x['options']=[]
    return {'items':items}
@app.post('/api/s2/custom-fields')
def s2_custom_add(body:CustomFieldIn,request:Request):
    u=current_user(request); require_role(u,'ADMIN')
    if body.field_type not in {'text','number','date','select'}: raise HTTPException(400,'Kiểu trường phải là text, number, date hoặc select.')
    if body.field_type=='select' and not body.options: raise HTTPException(400,'Trường danh sách chọn phải có lựa chọn.')
    with db_cursor(commit=True) as cur:
        cur.execute('INSERT INTO custom_fields(entity_type,field_name,field_type,is_required,options_json,sort_order) OUTPUT INSERTED.id VALUES(?,?,?,?,?,?)',body.entity_type,body.field_name,body.field_type,1 if body.is_required else 0,json.dumps(body.options,ensure_ascii=False),body.sort_order); cid=cur.fetchone()[0]; _audit(cur,u['id'],'created','CustomField',cid,'field_name',None,body.field_name)
    return {'message':'Đã thêm trường tùy chỉnh.','id':cid}

class StageIn(BaseModel): name:str; sort_order:int; win_probability:int; exit_condition:Optional[str]=None; is_active:bool=True
@app.get('/api/s2/pipeline-stages')
def s2_stages(request:Request):
    current_user(request)
    with db_cursor() as cur: cur.execute('SELECT * FROM pipeline_stages ORDER BY sort_order,id'); return {'items':rows_to_dicts(cur,cur.fetchall())}
@app.post('/api/s2/pipeline-stages')
def s2_stage_add(body:StageIn,request:Request):
    u=current_user(request); require_role(u,'DIRECTOR','ADMIN')
    if not 0<=body.win_probability<=100: raise HTTPException(400,'Xác suất thắng phải từ 0 đến 100%.')
    with db_cursor(commit=True) as cur:
        cur.execute('INSERT INTO pipeline_stages(name,sort_order,win_probability,exit_condition,is_active) OUTPUT INSERTED.id VALUES(?,?,?,?,?)',body.name,body.sort_order,body.win_probability,body.exit_condition,1 if body.is_active else 0); sid=cur.fetchone()[0]; _audit(cur,u['id'],'created','PipelineStage',sid,'name',None,body.name)
    return {'message':'Đã thêm giai đoạn pipeline.','id':sid}

class ReasonIn(BaseModel): result_type:str; reason:str; sort_order:int=0; is_active:bool=True
@app.get('/api/s2/win-loss-reasons')
def s2_reasons(request:Request):
    current_user(request)
    with db_cursor() as cur: cur.execute('SELECT * FROM win_loss_reasons ORDER BY result_type,sort_order,id'); return {'items':rows_to_dicts(cur,cur.fetchall())}
@app.post('/api/s2/win-loss-reasons')
def s2_reason_add(body:ReasonIn,request:Request):
    u=current_user(request); require_role(u,'DIRECTOR','ADMIN')
    if body.result_type not in {'Thắng','Thua'}: raise HTTPException(400,'Loại kết quả phải là Thắng hoặc Thua.')
    with db_cursor(commit=True) as cur:
        cur.execute('INSERT INTO win_loss_reasons(result_type,reason,sort_order,is_active) OUTPUT INSERTED.id VALUES(?,?,?,?)',body.result_type,body.reason,body.sort_order,1 if body.is_active else 0); rid=cur.fetchone()[0]; _audit(cur,u['id'],'created','WinLossReason',rid,'reason',None,body.reason)
    return {'message':'Đã thêm lý do thắng/thua.','id':rid}

# ---- S3-01 customer company profile + S3-07 search ----
class CustomerCompanyIn(BaseModel):
    name:str; tax_code:Optional[str]=None; industry:Optional[str]=None; company_size:Optional[str]=None; website:Optional[str]=None; address:Optional[str]=None; phone:Optional[str]=None; email:Optional[str]=None; status:str='Tiềm năng'; owner_id:Optional[int]=None; parent_customer_id:Optional[int]=None

@app.get('/api/s3/customers')
def s3_customers(request:Request,q:str='',status:str='',industry:str='',company_size:str='',region:str='',owner_id:int|None=None):
    u=current_user(request); clause,args=_customer_scope(u,'c'); wh=[clause]; params=list(args)
    if q: wh.append('(c.name LIKE ? OR c.tax_code LIKE ? OR EXISTS(SELECT 1 FROM contacts ct WHERE ct.customer_id=c.id AND ct.phone LIKE ?))'); s=f'%{q}%'; params += [s,s,s]
    if status: wh.append('c.status=?'); params.append(status)
    if industry: wh.append('c.industry=?'); params.append(industry)
    if company_size: wh.append('c.company_size=?'); params.append(company_size)
    if owner_id: wh.append('c.owner_id=?'); params.append(owner_id)
    if region: wh.append('g.region=?'); params.append(region)
    with db_cursor() as cur:
        cur.execute(f'''SELECT c.*,u2.full_name owner_name,g.name group_name,g.region,p.name parent_name FROM customers c JOIN users u2 ON u2.id=c.owner_id LEFT JOIN business_groups g ON g.id=c.business_group_id LEFT JOIN customers p ON p.id=c.parent_customer_id WHERE {' AND '.join(wh)} ORDER BY c.id DESC''',*params)
        return {'items':rows_to_dicts(cur,cur.fetchall()),'scope':u['data_scope']}

@app.post('/api/s3/customers')
def s3_customer_add(body:CustomerCompanyIn,request:Request):
    u=current_user(request); owner=body.owner_id or u['id']
    if body.status not in {'Tiềm năng','Đang giao dịch','Khách hàng','Ngừng hợp tác'}: raise HTTPException(400,'Trạng thái khách hàng không hợp lệ.')
    with db_cursor(commit=True) as cur:
        if body.tax_code:
            cur.execute('SELECT 1 FROM customers WHERE tax_code=?',body.tax_code)
            if cur.fetchone(): raise HTTPException(409,'Mã số thuế đã tồn tại.')
        cur.execute('SELECT business_group_id FROM users WHERE id=?',owner); rr=cur.fetchone(); gid=rr[0] if rr else u.get('business_group_id')
        cur.execute('''INSERT INTO customers(name,phone,email,owner_id,business_group_id,tax_code,industry,company_size,website,address,status,parent_customer_id,last_interaction_at) OUTPUT INSERTED.id VALUES(?,?,?,?,?,?,?,?,?,?,?,?,SYSUTCDATETIME())''',body.name,body.phone,body.email,owner,gid,body.tax_code,body.industry,body.company_size,body.website,body.address,body.status,body.parent_customer_id); cid=cur.fetchone()[0]; _audit(cur,u['id'],'created','Customer',cid,'name',None,body.name)
    return {'message':'Đã tạo hồ sơ khách hàng doanh nghiệp.','id':cid}

@app.put('/api/s3/customers/{cid:int}')
def s3_customer_update(cid:int,body:CustomerCompanyIn,request:Request):
    u=current_user(request)
    with db_cursor(commit=True) as cur:
        ensure_record_in_scope(cur,u,'customers',cid); cur.execute('SELECT * FROM customers WHERE id=?',cid); old=row_to_dict(cur,cur.fetchone())
        if body.tax_code:
            cur.execute('SELECT 1 FROM customers WHERE tax_code=? AND id<>?',body.tax_code,cid)
            if cur.fetchone(): raise HTTPException(409,'Mã số thuế đã tồn tại.')
        owner=body.owner_id or old['owner_id']; cur.execute('SELECT business_group_id FROM users WHERE id=?',owner); rr=cur.fetchone(); gid=rr[0] if rr else old['business_group_id']
        cur.execute('''UPDATE customers SET name=?,phone=?,email=?,owner_id=?,business_group_id=?,tax_code=?,industry=?,company_size=?,website=?,address=?,status=?,parent_customer_id=? WHERE id=?''',body.name,body.phone,body.email,owner,gid,body.tax_code,body.industry,body.company_size,body.website,body.address,body.status,body.parent_customer_id,cid)
        for f in ['status','owner_id','parent_customer_id']:
            nv=owner if f=='owner_id' else getattr(body,f)
            if old.get(f)!=nv: _audit(cur,u['id'],'updated','Customer',cid,f,old.get(f),nv)
    return {'message':'Đã cập nhật khách hàng.'}

# ---- S3-02 contacts ----
class ContactIn(BaseModel): full_name:str; title:Optional[str]=None; email:Optional[str]=None; phone:Optional[str]=None; buying_role:Optional[str]=None
@app.get('/api/s3/customers/{cid:int}/contacts')
def s3_contacts(cid:int,request:Request):
    u=current_user(request)
    with db_cursor() as cur:
        ensure_record_in_scope(cur,u,'customers',cid); cur.execute('SELECT * FROM contacts WHERE customer_id=? ORDER BY is_primary DESC,id',cid); return {'items':rows_to_dicts(cur,cur.fetchall())}
@app.post('/api/s3/customers/{cid:int}/contacts')
def s3_contact_add(cid:int,body:ContactIn,request:Request):
    u=current_user(request)
    if body.phone and not _PHONE_VN.fullmatch(body.phone): raise HTTPException(400,'Số điện thoại Việt Nam không hợp lệ.')
    if body.buying_role and body.buying_role not in {'Người quyết định','Người ảnh hưởng','Người dùng cuối','Người cản trở'}: raise HTTPException(400,'Vai trò quyết định mua không hợp lệ.')
    with db_cursor(commit=True) as cur:
        ensure_record_in_scope(cur,u,'customers',cid); cur.execute('INSERT INTO contacts(customer_id,full_name,title,email,phone,buying_role) OUTPUT INSERTED.id VALUES(?,?,?,?,?,?)',cid,body.full_name,body.title,body.email,body.phone,body.buying_role); contact_id=cur.fetchone()[0]; cur.execute('SELECT name FROM customers WHERE id=?',cid); company=cur.fetchone()[0]; cur.execute('INSERT INTO contact_company_history(contact_id,customer_id,company_name) VALUES(?,?,?)',contact_id,cid,company); _audit(cur,u['id'],'created','Contact',contact_id,'full_name',None,body.full_name)
    return {'message':'Đã thêm người liên hệ.','id':contact_id}
@app.put('/api/s3/contacts/{contact_id}/primary')
def s3_contact_primary(contact_id:int,request:Request):
    u=current_user(request)
    with db_cursor(commit=True) as cur:
        cur.execute('SELECT customer_id FROM contacts WHERE id=?',contact_id); r=cur.fetchone()
        if not r: raise HTTPException(404,'Không tìm thấy người liên hệ.')
        ensure_record_in_scope(cur,u,'customers',r[0]); cur.execute('UPDATE contacts SET is_primary=0 WHERE customer_id=?',r[0]); cur.execute('UPDATE contacts SET is_primary=1 WHERE id=?',contact_id); _audit(cur,u['id'],'updated','Contact',contact_id,'is_primary',0,1)
    return {'message':'Đã đặt làm đầu mối chính.'}
class ContactMoveIn(BaseModel): customer_id:int
@app.put('/api/s3/contacts/{contact_id}/move')
def s3_contact_move(contact_id:int,body:ContactMoveIn,request:Request):
    u=current_user(request)
    with db_cursor(commit=True) as cur:
        cur.execute('SELECT customer_id FROM contacts WHERE id=?',contact_id); r=cur.fetchone()
        if not r: raise HTTPException(404,'Không tìm thấy người liên hệ.')
        ensure_record_in_scope(cur,u,'customers',r[0]); ensure_record_in_scope(cur,u,'customers',body.customer_id)
        cur.execute('UPDATE contact_company_history SET to_at=SYSUTCDATETIME() WHERE contact_id=? AND to_at IS NULL',contact_id); cur.execute('SELECT name FROM customers WHERE id=?',body.customer_id); company=cur.fetchone()[0]; cur.execute('INSERT INTO contact_company_history(contact_id,customer_id,company_name) VALUES(?,?,?)',contact_id,body.customer_id,company); cur.execute('UPDATE contacts SET customer_id=?,is_primary=0 WHERE id=?',body.customer_id,contact_id); _audit(cur,u['id'],'updated','Contact',contact_id,'customer_id',r[0],body.customer_id)
    return {'message':'Đã chuyển công ty và giữ nguyên lịch sử.'}

# ---- S3-03 360 + activity/opportunity ----
class CustomerOpportunityIn(BaseModel): title:str; value:float=0; stage:str='Tiếp cận'; result_type:Optional[str]=None; result_reason:Optional[str]=None; competitor:Optional[str]=None
@app.post('/api/s3/customers/{cid:int}/opportunities')
def s3_opp_add(cid:int,body:CustomerOpportunityIn,request:Request):
    u=current_user(request)
    with db_cursor(commit=True) as cur:
        ensure_record_in_scope(cur,u,'customers',cid); prob_map={'Tiếp cận':10,'Xác định nhu cầu':25,'Đề xuất giải pháp':45,'Báo giá':65,'Đàm phán':80,'Chốt':100}; prob=prob_map.get(body.stage,20)
        cur.execute('SELECT owner_id,business_group_id FROM customers WHERE id=?',cid); own=cur.fetchone(); cur.execute('INSERT INTO opportunities(title,value,stage,owner_id,business_group_id,customer_id,win_probability,result_type,result_reason,competitor) OUTPUT INSERTED.id VALUES(?,?,?,?,?,?,?,?,?,?)',body.title,body.value,body.stage,own[0],own[1],cid,prob,body.result_type,body.result_reason,body.competitor); oid=cur.fetchone()[0]; _audit(cur,u['id'],'created','Opportunity',oid,'title',None,body.title)
    return {'message':'Đã tạo cơ hội.','id':oid}
class ActivityIn(BaseModel): subject:str; activity_type:str='Ghi chú'
@app.post('/api/s3/customers/{cid:int}/activities')
def s3_activity_add(cid:int,body:ActivityIn,request:Request):
    u=current_user(request)
    with db_cursor(commit=True) as cur:
        ensure_record_in_scope(cur,u,'customers',cid); cur.execute('SELECT owner_id,business_group_id FROM customers WHERE id=?',cid); own=cur.fetchone(); cur.execute('INSERT INTO activities(subject,activity_type,owner_id,business_group_id,customer_id) OUTPUT INSERTED.id VALUES(?,?,?,?,?)',body.subject,body.activity_type,own[0],own[1],cid); aid=cur.fetchone()[0]; cur.execute('UPDATE customers SET last_interaction_at=SYSUTCDATETIME() WHERE id=?',cid)
    return {'message':'Đã ghi hoạt động.','id':aid}
@app.post('/api/s3/customers/{cid:int}/files')
async def s3_customer_file_upload(cid:int, request:Request, file:UploadFile=File(...)):
    u=current_user(request)
    raw=await file.read()
    if not raw:
        raise HTTPException(400,'Tệp rỗng.')
    if len(raw)>10*1024*1024:
        raise HTTPException(400,'Tệp vượt quá giới hạn 10MB.')
    safe_name=_Path(file.filename or 'tep-dinh-kem').name
    with db_cursor() as cur:
        ensure_record_in_scope(cur,u,'customers',cid)
    folder=_UPLOAD_DIR / 'customer_files'
    folder.mkdir(parents=True, exist_ok=True)
    stored=folder / f"kh{cid}_{int(datetime.now().timestamp()*1000)}_{safe_name}"
    stored.write_bytes(raw)
    with db_cursor(commit=True) as cur:
        cur.execute('INSERT INTO customer_files(customer_id,file_name,file_path,file_size,uploader_id) OUTPUT INSERTED.id VALUES(?,?,?,?,?)',cid,safe_name,str(stored),len(raw),u['id'])
        fid=cur.fetchone()[0]
        _audit(cur,u['id'],'uploaded','CustomerFile',fid,'file_name',None,safe_name)
    return {'message':'Đã tải tệp đính kèm.','id':fid,'file_name':safe_name}

@app.get('/api/s3/customer-files/{file_id:int}')
def s3_customer_file_download(file_id:int,request:Request):
    u=current_user(request)
    with db_cursor() as cur:
        cur.execute('SELECT * FROM customer_files WHERE id=?',file_id)
        f=row_to_dict(cur,cur.fetchone())
        if not f:
            raise HTTPException(404,'Không tìm thấy tệp đính kèm.')
        ensure_record_in_scope(cur,u,'customers',f['customer_id'])
    if not os.path.exists(f['file_path']):
        raise HTTPException(404,'Tệp không còn tồn tại trên máy chủ.')
    return FileResponse(f['file_path'],filename=f['file_name'])

@app.get('/api/s3/customers/{cid:int}/360')
def s3_360(cid:int,request:Request):
    import time as _time; t=_time.perf_counter(); u=current_user(request)
    with db_cursor() as cur:
        ensure_record_in_scope(cur,u,'customers',cid); cur.execute('SELECT c.*,u2.full_name owner_name,g.name group_name,p.name parent_name FROM customers c JOIN users u2 ON u2.id=c.owner_id LEFT JOIN business_groups g ON g.id=c.business_group_id LEFT JOIN customers p ON p.id=c.parent_customer_id WHERE c.id=?',cid); customer=row_to_dict(cur,cur.fetchone())
        cur.execute('SELECT * FROM contacts WHERE customer_id=? ORDER BY is_primary DESC,id',cid); contacts=rows_to_dicts(cur,cur.fetchall())
        cur.execute('SELECT * FROM opportunities WHERE customer_id=? ORDER BY id DESC',cid); opps=rows_to_dicts(cur,cur.fetchall())
        cur.execute('SELECT TOP 500 * FROM activities WHERE customer_id=? ORDER BY id DESC',cid); acts=rows_to_dicts(cur,cur.fetchall())
        cur.execute('SELECT * FROM customer_files WHERE customer_id=? ORDER BY id DESC',cid); files=rows_to_dicts(cur,cur.fetchall())
        signed=sum(float(x['value'] or 0) for x in opps if x.get('result_type')=='Thắng' or x.get('stage')=='Chốt'); opened=sum(float(x['value'] or 0) for x in opps if not x.get('result_type'))
    elapsed=round(_time.perf_counter()-t,4)
    return {'customer':customer,'contacts':contacts,'opportunities':opps,'activities':acts,'files':files,'total_signed':signed,'total_open':opened,'load_seconds':elapsed,'performance_pass':elapsed<1.5}

# ---- S3-04 duplicates/merge ----
@app.get('/api/s3/customers/duplicates')
def s3_duplicates(request:Request):
    u=current_user(request); require_role(u,'TEAM_LEADER','DIRECTOR','ADMIN')
    with db_cursor() as cur:
        clause,args=_customer_scope(u,'c'); cur.execute(f'''SELECT c.id,c.name,c.tax_code,c.website,c.phone,c.email,c.owner_id,u.full_name owner_name FROM customers c JOIN users u ON u.id=c.owner_id WHERE {clause} ORDER BY c.id''',*args); items=rows_to_dicts(cur,cur.fetchall())
    pairs=[]
    for i,a in enumerate(items):
        for b in items[i+1:]:
            reasons=[]
            if a.get('tax_code') and a['tax_code']==b.get('tax_code'): reasons.append('trùng mã số thuế')
            if a.get('website') and a['website'].lower()==str(b.get('website') or '').lower(): reasons.append('trùng website')
            na=''.join(ch for ch in a['name'].lower() if ch.isalnum()); nb=''.join(ch for ch in b['name'].lower() if ch.isalnum())
            if na and nb and (na in nb or nb in na): reasons.append('tên công ty gần giống')
            if reasons: pairs.append({'a':a,'b':b,'reasons':reasons})
    return {'items':pairs}
class MergeIn(BaseModel): keep_id:int; merge_id:int
@app.post('/api/s3/customers/merge')
def s3_merge(body:MergeIn,request:Request):
    u=current_user(request); require_role(u,'TEAM_LEADER','DIRECTOR','ADMIN')
    if body.keep_id==body.merge_id: raise HTTPException(400,'Hai khách hàng phải khác nhau.')
    with db_cursor(commit=True) as cur:
        ensure_record_in_scope(cur,u,'customers',body.keep_id); ensure_record_in_scope(cur,u,'customers',body.merge_id)
        cur.execute('UPDATE contacts SET customer_id=? WHERE customer_id=?',body.keep_id,body.merge_id); cur.execute('UPDATE opportunities SET customer_id=? WHERE customer_id=?',body.keep_id,body.merge_id); cur.execute('UPDATE activities SET customer_id=? WHERE customer_id=?',body.keep_id,body.merge_id); cur.execute('UPDATE customer_files SET customer_id=? WHERE customer_id=?',body.keep_id,body.merge_id); cur.execute('UPDATE support_tickets SET customer_id=? WHERE customer_id=?',body.keep_id,body.merge_id); cur.execute('UPDATE customers SET parent_customer_id=? WHERE parent_customer_id=?',body.keep_id,body.merge_id); cur.execute('DELETE FROM customers WHERE id=?',body.merge_id); _audit(cur,u['id'],'merged','Customer',body.keep_id,'merged_from',body.merge_id,body.keep_id)
    return {'message':'Đã gộp khách hàng và giữ lại người liên hệ, cơ hội, hoạt động và tệp đính kèm.'}

# ---- S3-05 hierarchy ----
@app.get('/api/s3/customers/{cid:int}/hierarchy')
def s3_hierarchy(cid:int,request:Request):
    u=current_user(request)
    with db_cursor() as cur:
        ensure_record_in_scope(cur,u,'customers',cid); cur.execute('SELECT id,name,parent_customer_id FROM customers WHERE id=?',cid); root=row_to_dict(cur,cur.fetchone()); cur.execute('SELECT id,name,status FROM customers WHERE parent_customer_id=? ORDER BY name',cid); children=rows_to_dicts(cur,cur.fetchall()); ids=[cid]+[x['id'] for x in children]
        total=0
        for x in ids:
            cur.execute("SELECT COALESCE(SUM(value),0) FROM opportunities WHERE customer_id=? AND (result_type=N'Thắng' OR stage=N'Chốt')",x); total += float(cur.fetchone()[0] or 0)
    return {'company':root,'children':children,'group_contract_value':total}

# ---- S3-06 customer import ----
@app.get('/api/s3/customers/import-template')
def s3_customer_template(request:Request):
    current_user(request); wb=Workbook(); ws=wb.active; ws.title='Khách hàng'; ws.append(['Tên công ty','Mã số thuế','Ngành nghề','Quy mô','Website','Địa chỉ','Điện thoại','Email','Trạng thái']); ws.append(['Công ty Mẫu','0101234567','Công nghệ','51-200','https://example.vn','Hà Nội','0912345678','contact@example.vn','Tiềm năng']); b=io.BytesIO(); wb.save(b); b.seek(0); return StreamingResponse(b,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':'attachment; filename=tep-mau-khach-hang.xlsx'})
def _parse_customer_xlsx(raw,u):
    from openpyxl import load_workbook
    try: ws=load_workbook(io.BytesIO(raw),data_only=True).active
    except: raise HTTPException(400,'Tệp Excel không hợp lệ.')
    out=[]
    with db_cursor() as cur: cur.execute("SELECT tax_code FROM customers WHERE tax_code IS NOT NULL"); existing={r[0] for r in cur.fetchall()}
    seen=set()
    for idx,r in enumerate(ws.iter_rows(min_row=2,values_only=True),start=2):
        if not any(x is not None and str(x).strip() for x in r): continue
        vals=[str(x or '').strip() for x in list(r)+[None]*9][:9]; name,tax,industry,size,website,address,phone,email,status=vals; errors=[]
        if not name: errors.append('Thiếu tên công ty')
        if tax and (tax in existing or tax in seen): errors.append('Mã số thuế trùng')
        if status not in {'Tiềm năng','Đang giao dịch','Khách hàng','Ngừng hợp tác'}: errors.append('Trạng thái không hợp lệ')
        seen.add(tax); out.append({'row':idx,'valid':not errors,'duplicate':bool(tax and tax in existing),'errors':errors,'data':{'name':name,'tax_code':tax or None,'industry':industry,'company_size':size,'website':website,'address':address,'phone':phone,'email':email,'status':status}})
    return out
@app.post('/api/s3/customers/import-preview')
async def s3_customer_preview(request:Request,file:UploadFile=File(...)):
    u=current_user(request); rows=_parse_customer_xlsx(await file.read(),u); return {'items':rows,'valid':sum(x['valid'] for x in rows),'invalid':sum(not x['valid'] for x in rows)}
@app.post('/api/s3/customers/import')
async def s3_customer_import(request:Request,file:UploadFile=File(...)):
    u=current_user(request); rows=_parse_customer_xlsx(await file.read(),u); created=[]; skipped=[]
    with db_cursor(commit=True) as cur:
        for x in rows:
            if not x['valid']: skipped.append(x); continue
            d=x['data']; cur.execute('INSERT INTO customers(name,phone,email,owner_id,business_group_id,tax_code,industry,company_size,website,address,status,last_interaction_at) OUTPUT INSERTED.id VALUES(?,?,?,?,?,?,?,?,?,?,?,SYSUTCDATETIME())',d['name'],d['phone'],d['email'],u['id'],u.get('business_group_id'),d['tax_code'],d['industry'],d['company_size'],d['website'],d['address'],d['status']); cid=cur.fetchone()[0]; created.append({'row':x['row'],'id':cid,'name':d['name']})
    return {'message':f'Đã nhập {len(created)} khách hàng, bỏ qua {len(skipped)} dòng lỗi/trùng.','created':created,'skipped':skipped}

# ---- S3-07 saved filters ----
class SavedFilterIn(BaseModel): name:str; filters:dict
@app.get('/api/s3/customer-filters')
def s3_saved_filters(request:Request):
    u=current_user(request)
    with db_cursor() as cur: cur.execute('SELECT id,name,filter_json,created_at FROM saved_customer_filters WHERE user_id=? ORDER BY id DESC',u['id']); items=rows_to_dicts(cur,cur.fetchall())
    for x in items:
        try:x['filters']=json.loads(x['filter_json'])
        except:x['filters']={}
    return {'items':items}
@app.post('/api/s3/customer-filters')
def s3_saved_filter_add(body:SavedFilterIn,request:Request):
    u=current_user(request)
    with db_cursor(commit=True) as cur: cur.execute('INSERT INTO saved_customer_filters(user_id,name,filter_json) OUTPUT INSERTED.id VALUES(?,?,?)',u['id'],body.name,json.dumps(body.filters,ensure_ascii=False)); fid=cur.fetchone()[0]
    return {'message':'Đã lưu bộ lọc.','id':fid}

# ---- S3-08 support/churn ----
class TicketIn(BaseModel): title:str; priority:str='Trung bình'; assignee_id:Optional[int]=None
@app.get('/api/s3/support')
def s3_support(request:Request):
    u=current_user(request); clause,args=_customer_scope(u,'c')
    with db_cursor() as cur:
        cur.execute(f'''SELECT t.id,t.customer_id,c.name customer_name,t.title,t.priority,t.status,t.assignee_id,u2.full_name assignee_name,t.created_at FROM support_tickets t JOIN customers c ON c.id=t.customer_id LEFT JOIN users u2 ON u2.id=t.assignee_id WHERE {clause} ORDER BY t.id DESC''',*args); items=rows_to_dicts(cur,cur.fetchall())
        cur.execute(f'''SELECT c.id,c.name,COUNT(CASE WHEN t.status<>N'Đã đóng' THEN 1 END) open_count FROM customers c LEFT JOIN support_tickets t ON t.customer_id=c.id WHERE {clause} GROUP BY c.id,c.name HAVING COUNT(CASE WHEN t.status<>N'Đã đóng' THEN 1 END)>=3''',*args); risks=rows_to_dicts(cur,cur.fetchall())
    return {'items':items,'churn_risks':risks}
@app.post('/api/s3/customers/{cid:int}/support')
def s3_ticket_add(cid:int,body:TicketIn,request:Request):
    u=current_user(request)
    with db_cursor(commit=True) as cur:
        ensure_record_in_scope(cur,u,'customers',cid); cur.execute('INSERT INTO support_tickets(customer_id,title,priority,assignee_id) OUTPUT INSERTED.id VALUES(?,?,?,?)',cid,body.title,body.priority,body.assignee_id or u['id']); tid=cur.fetchone()[0]
    return {'message':'Đã ghi nhận yêu cầu hỗ trợ.','id':tid}

@app.put('/api/s3/support/{ticket_id:int}/close')
def s3_support_close(ticket_id:int,request:Request):
    u=current_user(request)
    with db_cursor(commit=True) as cur:
        cur.execute('SELECT t.id,t.customer_id,t.status FROM support_tickets t WHERE t.id=?',ticket_id)
        t=row_to_dict(cur,cur.fetchone())
        if not t:
            raise HTTPException(404,'Không tìm thấy yêu cầu hỗ trợ.')
        ensure_record_in_scope(cur,u,'customers',t['customer_id'])
        cur.execute("UPDATE support_tickets SET status=N'Đã đóng',closed_at=SYSUTCDATETIME() WHERE id=?",ticket_id)
    return {'message':'Đã đóng yêu cầu hỗ trợ.'}

# ---- S3-09 periodic care ----
@app.get('/api/s3/care-list')
def s3_care_list(request:Request,days:int=30):
    u=current_user(request); clause,args=_customer_scope(u,'c'); days=max(1,min(days,3650))
    with db_cursor() as cur:
        cur.execute(f'''SELECT c.id,c.name,c.status,c.last_interaction_at,u2.full_name owner_name,COALESCE(SUM(CASE WHEN o.result_type=N'Thắng' OR o.stage=N'Chốt' THEN o.value ELSE 0 END),0) contract_value FROM customers c JOIN users u2 ON u2.id=c.owner_id LEFT JOIN opportunities o ON o.customer_id=c.id WHERE {clause} AND (c.last_interaction_at IS NULL OR c.last_interaction_at<DATEADD(day,-?,SYSUTCDATETIME())) GROUP BY c.id,c.name,c.status,c.last_interaction_at,u2.full_name ORDER BY contract_value DESC,c.id''',*(args+[days])); return {'items':rows_to_dicts(cur,cur.fetchall()),'days':days}
class CareDoneIn(BaseModel): note:Optional[str]=None
@app.post('/api/s3/customers/{cid:int}/care-done')
def s3_care_done(cid:int,body:CareDoneIn,request:Request):
    u=current_user(request)
    with db_cursor(commit=True) as cur:
        ensure_record_in_scope(cur,u,'customers',cid); cur.execute('UPDATE customers SET last_interaction_at=SYSUTCDATETIME() WHERE id=?',cid); cur.execute('INSERT INTO care_logs(customer_id,user_id,note) VALUES(?,?,?)',cid,u['id'],body.note)
    return {'message':'Đã đánh dấu liên hệ hôm nay.'}

@app.post('/api/users/{user_id}/unlock')
def unlock_user(user_id:int,request:Request):
    actor=current_user(request); require_role(actor,'ADMIN')
    with db_cursor(commit=True) as cur:
        cur.execute("UPDATE users SET is_active=1,status=N'active',failed_login_attempts=0,locked_until=NULL,updated_at=? WHERE id=?",utcnow(),user_id)
        if cur.rowcount==0: raise HTTPException(404,'Không tìm thấy tài khoản.')
        _audit(cur,actor['id'],'unlocked','User',user_id,'status','locked','active')
    return {'message':'Đã mở khóa tài khoản. Người dùng có thể đăng nhập lại.'}
